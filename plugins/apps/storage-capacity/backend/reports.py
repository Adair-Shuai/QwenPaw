# -*- coding: utf-8 -*-
"""One report model for preview, Word, Markdown and reproducible exports."""

from __future__ import annotations

import html
import io
import json
from zipfile import ZipFile, ZIP_DEFLATED

from .catalog import METHODS, TERMS
from .workbench import METRICS

DIAGNOSTICS = {
    "start_pz": "起始 p/Z（MPa）",
    "end_pz": "末期 p/Z（MPa）",
    "inverse_withdrawal_fraction": "撤采比例倒数",
    "intercept": "回归截距",
    "slope": "回归斜率",
    "r_squared": "拟合优度 R²",
    "rmse": "均方根误差",
    "mae": "平均绝对误差",
    "n": "样本数",
    "metric": "拟合指标",
    "unit": "拟合单位",
    "schedule_integral": "采气曲线积分（亿m³）",
    "net_water_reservoir_m3": "净水侵量（地层m³）",
    "bg_depletion": "体积系数变化",
}
STATES = {
    "start": "评估起始状态",
    "upper_limit": "上限压力状态",
    "lower_limit": "下限压力状态",
    "pressure_window": "规定压力窗口",
    "start_to_lower": "起始至下限压力区间",
    "withdrawal_period": "规定采气期",
    "deliverability_condition": "规定井底与外输工况",
}


def condition_text(c):
    labels = {
        "water_drive": "水体影响",
        "fluid": "流体与相态",
        "pressure_representative": "压力／体积代表性",
        "closed_during_segment": "本段近似定容且无注气",
        "interference": "井间干扰",
        "water_volume_verified": "水侵参数核实",
    }
    choices = {
        "unknown": "尚未判断",
        "none": "无",
        "weak": "弱，已核实",
        "moderate": "中等",
        "strong": "强",
        "dry_gas": "单气相为主",
        "stable_gas": "油环／凝析影响已论证可忽略",
        "oil": "油环影响显著",
        "condensate": "凝析影响显著",
        True: "已确认",
        False: "未确认",
    }
    return "；".join(
        f"{labels.get(k, k)}：{choices.get(v, str(v))}" for k, v in c.items()
    )


def fmt(value):
    if value is None:
        return "未计算"
    return (
        f"{value:,.4f}".rstrip("0").rstrip(".")
        if isinstance(value, (int, float))
        else str(value)
    )


def sections(evaluation, review=None):  # pylint: disable=too-many-branches
    case = evaluation["case"]
    data = []

    def section(title, paragraphs=None, headers=None, rows=None):
        data.append(
            {
                "title": title,
                "paragraphs": paragraphs or [],
                "headers": headers or [],
                "rows": rows or [],
            },
        )

    section(
        "评估结论与范围",
        [
            evaluation["summary"],
            f"评价周期：{case.get('cycle_id') or '未填写'}。范围："
            f"{'、'.join(u['name'] for u in evaluation['units']) or '未划分'}。",
            f"起始状态：{case.get('start_state') or '未填写'}；"
            f"末期状态：{case.get('end_state') or '未填写'}。",
            "压力口径："
            + {
                "absolute": "绝对压力",
                "apparent_formation": "视地层压力",
                "report_defined": "报告定义口径",
            }.get(case.get("pressure_basis"), "未确认")
            + "。气量统一为亿标准立方米，日采气量为万标准立方米/日。",
            (
                "气量标准状态：沿用同一资料的报告口径作条件复算，标准温压待核实。"
                if case.get("standard", {}).get("mode") == "reported"
                else f"标准状态：{case.get('standard', {}).get('temperature')} K，"
                f"{case.get('standard', {}).get('pressure')} MPa，"
                f"Zs={case.get('standard', {}).get('z')}。"
            )
            + f"依据：{case.get('standard', {}).get('source') or '未填写'}。",
            "结果状态：计算建议值，待复核。尚无正式签审或批准依据。",
        ],
    )
    section(
        "主要指标",
        headers=["指标", "结果", "单位", "依据／状态"],
        rows=[
            [
                m["label"],
                fmt(m["value"]),
                m["unit"],
                m["source"] + "；" + STATES.get(m["state"], m["state"]),
            ]
            for m in evaluation["metrics"]
        ],
    )
    section(
        "资料基础",
        headers=["资料", "内容摘要", "文件校验"],
        rows=[
            [
                d["name"],
                f"{len(d.get('tables', []))} 张表，"
                f"{len(d.get('pages', []))} 个文本段；"
                + {"extracted": "已读取", "manual_review": "需核对原件"}.get(
                    d.get("status"),
                    "已登记",
                ),
                d.get("sha256", "未记录"),
            ]
            for d in case.get("documents", [])
        ],
    )
    section(
        "质量与口径检查",
        headers=["检查项", "状态", "检查记录"],
        rows=[
            [
                c["label"],
                {"passed": "通过", "warning": "需复核", "error": "待补充"}.get(
                    c["status"],
                    c["status"],
                ),
                c["detail"],
            ]
            for c in evaluation["checks"]
        ],
    )
    for unit, item in zip(case.get("units", []), evaluation["units"]):
        method = next(
            (m for m in METHODS if m["id"] == item["method"]),
            {
                "name": str(item["method"]),
                "formula": "",
                "condition": "",
                "reference": "",
            },
        )
        section(
            f"{item['name']} 方法与结果",
            [
                "采用方法：" + method["name"],
                "输入依据：" + (unit.get("source") or "未填写"),
                "计算关系：" + method["formula"],
                "适用条件：" + method["condition"],
                "方法出处：" + method["reference"],
                "条件声明：" + condition_text(unit.get("conditions", {})),
                *item.get("errors", []),
                *item.get("warnings", []),
                *item.get("steps", []),
            ],
            headers=["指标", "结果", "单位"],
            rows=[
                [
                    METRICS.get(k, k),
                    fmt(v),
                    "万m³/d" if k == "peak_rate" else "亿m³",
                ]
                for k, v in item.get("values", {}).items()
            ],
        )
        p = item.get("resolved_parameters", unit.get("parameters", {}))
        from .catalog import FIELDS

        section(
            f"{item['name']} 输入参数",
            headers=["参数", "取值", "单位", "定位"],
            rows=[
                [
                    FIELDS.get(k, (k, ""))[0],
                    (
                        format(v, ".10g")
                        if isinstance(v, (int, float))
                        else str(v)
                    ),
                    FIELDS.get(k, ("", ""))[1],
                    (unit.get("evidence", {}).get(k) or {}).get(
                        "location",
                        unit.get("source", "用户填写"),
                    ),
                ]
                for k, v in p.items()
                if v not in (None, "")
            ],
        )
        if item.get("diagnostics"):
            section(
                f"{item['name']} 计算诊断",
                headers=["项目", "结果"],
                rows=[
                    [DIAGNOSTICS.get(k, k), fmt(v)]
                    for k, v in item["diagnostics"].items()
                    if k != "residuals"
                ],
            )
        if item.get("plot"):
            section(
                f"{item['name']} 回归数据与残差",
                headers=["累采量（亿m³）", "实测 p/Z", "拟合 p/Z", "残差"],
                rows=[
                    [
                        fmt(r["x"]),
                        fmt(r["y"]),
                        fmt(r["fitted"]),
                        fmt(r["y"] - r["fitted"]),
                    ]
                    for r in item["plot"]
                ],
            )
        for key, label in (
            ("start_pressures", "起始状态测压"),
            ("end_pressures", "末期状态测压"),
        ):
            if unit.get(key):
                section(
                    f"{item['name']} {label}",
                    [
                        f"权重依据：{unit.get('pressure_weight_basis', '')}；"
                        f"{unit.get('pressure_source', '')}",
                    ],
                    ["井名", "压力（MPa）", "权重"],
                    [
                        [
                            r.get("name", ""),
                            fmt(r.get("pressure")),
                            fmt(r.get("weight")),
                        ]
                        for r in unit[key]
                    ],
                )
        if unit.get("schedule"):
            section(
                f"{item['name']} 采气预测",
                [unit.get("schedule_source", "")],
                ["区间结束日", "区间平均日量（万m³/d）"],
                [[fmt(r["day"]), fmt(r["rate"])] for r in unit["schedule"]],
            )
        if item.get("well_results"):
            section(
                f"{item['name']} 单井产能筛选",
                headers=["井名", "产能（万m³/d）"],
                rows=[
                    [r["name"], fmt(r["rate"])] for r in item["well_results"]
                ],
            )
        if item.get("sensitivity"):
            s = item["sensitivity"]
            section(
                f"{item['name']} 敏感性",
                [
                    f"输入扰动 ±{s['percent']:g}%，{s['label']}。空值表示扰动后不满足计算条件。",
                ],
                ["输入", "负向扰动结果", "正向扰动结果"],
                [
                    [r["parameter"], fmt(r["minus"]), fmt(r["plus"])]
                    for r in s["rows"]
                ],
            )
    section(
        "交叉校验",
        [
            "比较阈值由项目指定，只作差异提示，不作为行业合格线。不同口径或不重叠性未确认的结果不直接平均。",
        ],
        ["指标", "本次计算", "对比值", "偏差 %", "对比依据"],
        [
            [
                METRICS.get(r["metric"], r["metric"]),
                fmt(r["calculated"]),
                fmt(r["value"]),
                fmt(r["difference_percent"]),
                r["source"],
            ]
            for r in evaluation.get("comparisons", [])
        ],
    )
    if case.get("history"):
        section(
            "多周期指标记录",
            [
                "以下为导入的历史成果，未重新计算。各周期压力、井网与运行约束可能变化，不能仅凭曲线偏移认定漏失。",
            ],
            ["周期", "账面库存", "有效库存", "综合调峰能力", "来源"],
            [
                [
                    r.get("cycle_id", ""),
                    fmt(r.get("book_inventory")),
                    fmt(r.get("effective_inventory")),
                    fmt(r.get("deliverable_swing")),
                    r.get("source", "未填写"),
                ]
                for r in case["history"]
            ],
        )
    section(
        "工程复核意见",
        [
            (
                review["text"]
                if review and review.get("status") == "completed"
                else "尚未完成智能体方法审查。算法结果及资料检查见前文。"
            ),
            case.get("review_notes") or "项目复核意见待填写。",
        ],
    )
    if evaluation.get("workflow"):
        status_names = {
            "calculated": "已计算",
            "blocked": "缺项阻断",
            "stale": "输入已变，旧结果不采用",
            "review": "待复核",
            "not_applicable": "不适用",
        }
        section(
            "流程节点与参数传递",
            [
                "各节点按自身输入独立运行；过期节点仅保留记录，不作为本次采用成果。",
            ],
            ["节点", "范围", "状态", "输入版本"],
            [
                [
                    r["node"],
                    (
                        "全项目"
                        if r["unit_index"] is None
                        else case["units"][r["unit_index"]]["name"]
                    ),
                    status_names.get(r["status"], r["status"]),
                    r["snapshot"]["input_hash"],
                ]
                for r in evaluation["workflow"]
            ],
        )
        for r in evaluation["workflow"]:
            if r["status"] == "stale":
                continue
            snapshot = r["snapshot"]
            section(
                r["node"] + " 节点记录",
                [
                    *snapshot.get("steps", []),
                    *snapshot.get("errors", []),
                    *snapshot.get("warnings", []),
                ],
            )
            if snapshot.get("values"):
                section(
                    r["node"] + " 节点结果",
                    headers=["指标", "取值"],
                    rows=[
                        [METRICS.get(k, k), fmt(v)]
                        for k, v in snapshot["values"].items()
                    ],
                )
            if snapshot.get("rows"):
                rows = snapshot["rows"]
                headers = list(rows[0])
                section(
                    r["node"] + " 过程数据",
                    (
                        [
                            "超过 30 行时，报告展示前 30 行，完整数据保留在复算包。",
                        ]
                        if len(rows) > 30
                        else []
                    ),
                    headers,
                    [[fmt(row.get(k)) for k in headers] for row in rows[:30]],
                )
    if case.get("field_sources"):
        section(
            "对话填写与取值来源",
            ["以下为已采用的字段来源，数值计算以计算节点记录为准。"],
            ["字段", "位置", "原文证据", "采用依据"],
            [
                [
                    path,
                    source.get("location", ""),
                    source.get("quote", ""),
                    source.get("reason", ""),
                ]
                for path, source in case["field_sources"].items()
            ],
        )
    section(
        "标准与术语依据",
        [
            f"{s['id']}《{s['name']}》：{s['basis']}"
            for s in evaluation["standards"]
        ]
        + [
            "标准引用与工程实现分开记录。未核实全文的条款号未写入本报告。",
        ],
        ["术语", "本报告口径"],
        [[name, definition] for _, name, definition in TERMS],
    )
    section(
        "计算记录",
        [
            f"计算引擎版本：{evaluation['engine_version']}。"
            f"计算时间：{evaluation['created_at']}。",
            f"输入 SHA-256：{evaluation['fingerprint']}",
            "复算包中的 JSON 保留输入、公式结果、质量检查、外部成果依据及版本。更新任何输入后应重新计算并生成报告。",
        ],
    )
    return data


def markdown(evaluation, review=None):
    title = evaluation["case"].get("case_name") or "储气库"
    lines = [f"# {title} 库容评估报告", ""]
    for s in sections(evaluation, review):
        lines.extend(["## " + s["title"], "", *s["paragraphs"], ""])
        if s["headers"] and s["rows"]:
            escape = lambda v: str(v).replace("|", "／").replace("\n", " ")
            lines.append("| " + " | ".join(s["headers"]) + " |")
            lines.append("| " + " | ".join("---" for _ in s["headers"]) + " |")
            lines.extend(
                "| " + " | ".join(escape(v) for v in r) + " |"
                for r in s["rows"]
            )
            lines.append("")
    return "\n".join(lines)


def html_report(evaluation, review=None):
    esc = lambda v: html.escape(str(v))
    content = [
        "<h1>"
        + esc(evaluation["case"].get("case_name", "储气库"))
        + " 库容评估报告</h1>",
    ]
    for s in sections(evaluation, review):
        content.append("<section><h2>" + esc(s["title"]) + "</h2>")
        content.extend("<p>" + esc(p) + "</p>" for p in s["paragraphs"])
        if s["rows"]:
            content.append(
                "<table><thead><tr>"
                + "".join("<th>" + esc(h) + "</th>" for h in s["headers"])
                + "</tr></thead><tbody>",
            )
            content.extend(
                "<tr>"
                + "".join("<td>" + esc(c) + "</td>" for c in row)
                + "</tr>"
                for row in s["rows"]
            )
            content.append("</tbody></table>")
        content.append("</section>")
    return (
        (
            '<!doctype html><html lang="zh-CN"><meta char'
            'set="utf-8"><title>库容评估报告</title><style>body'
            "{max-width:920px;margin:36px auto;padding:0 "
            '24px;font:14px/1.8 "Microsoft YaHei",sans-se'
            "rif;color:#182033}h1{font-size:25px}h2{font-"
            "size:19px;margin-top:32px}p{white-space:pre-"
            "wrap;overflow-wrap:anywhere}table{border-col"
            "lapse:collapse;width:100%;font-size:12px;tab"
            "le-layout:fixed}td,th{border:1px solid #d8dc"
            "e3;padding:8px;text-align:left;overflow-wrap"
            ":anywhere}th{background:#f3f5f8}thead{displa"
            "y:table-header-group}tr{break-inside:avoid}@"
            "media print{body{margin:0;max-width:none}h2{"
            "break-after:avoid}}@page{size:A4;margin:18mm"
            "}</style>"
        )
        + "".join(content)
        + "</html>"
    )


def docx_report(evaluation, review=None):
    # XML operations preserve table layout and repeat headers in Word.
    # pylint: disable=protected-access
    # pylint: disable=too-many-branches,too-many-statements
    # pylint: disable=too-many-nested-blocks
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in [
        ("Normal", 10),
        ("Title", 18),
        ("Heading 1", 14),
        ("Heading 2", 12),
    ]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(
            qn("w:eastAsia"),
            "Microsoft YaHei",
        )
        style.paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.25
    # The built-in template can carry a blue Title border; remove inherited
    # rules.
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    doc.add_paragraph(
        (evaluation["case"].get("case_name") or "储气库") + "\n库容评估报告",
        "Title",
    )
    for s in sections(evaluation, review):
        doc.add_heading(s["title"], level=1)
        for p in s["paragraphs"]:
            doc.add_paragraph(p)
        if s["rows"]:
            table = doc.add_table(rows=1, cols=len(s["headers"]))
            table.style = "Table Grid"
            table.autofit = False
            widths = {
                2: [4.5, 12.5],
                3: [5, 4, 8],
                4: [4, 2.7, 2, 8.3],
                5: [3.2, 2.3, 2.3, 2.2, 7],
            }.get(
                len(s["headers"]),
                [17 / len(s["headers"])] * len(s["headers"]),
            )
            if s["title"] == "质量与口径检查":
                widths = [3.1, 1.6, 12.3]
            if s["title"] == "资料基础":
                widths = [5.5, 4, 7.5]
            if s["title"].endswith("输入参数"):
                widths = [5, 3, 2.5, 6.5]
            for column, width in zip(table.columns, widths):
                column.width = Cm(width)
            for cell, header in zip(table.rows[0].cells, s["headers"]):
                cell.text = header
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EEF1F5")
                cell._tc.get_or_add_tcPr().append(
                    shading,
                )  # pylint: disable=protected-access
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(
                repeat,
            )  # pylint: disable=protected-access
            for row in s["rows"]:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = str(value)
                no_split = OxmlElement("w:cantSplit")
                cells[0]._tc.getparent().get_or_add_trPr().append(
                    no_split,
                )  # pylint: disable=protected-access
            for index, row in enumerate(table.rows):
                for cell, width in zip(row.cells, widths):
                    cell.width = Cm(width)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_before = Pt(2)
                        paragraph.paragraph_format.space_after = Pt(2)
                        paragraph.paragraph_format.line_spacing = 1.1
                        paragraph.paragraph_format.keep_with_next = index == 0
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
    doc.core_properties.author = "UGSci"
    doc.core_properties.title = "库容评估报告"
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def bundle(evaluation, review=None):
    stream = io.BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as z:
        z.writestr("库容评估报告.html", html_report(evaluation, review))
        z.writestr("库容评估报告.md", markdown(evaluation, review))
        z.writestr(
            "evaluation.json",
            json.dumps(
                evaluation,
                ensure_ascii=False,
                indent=2,
                allow_nan=False,
            ),
        )
        z.writestr(
            "case.json",
            json.dumps(
                evaluation["case"],
                ensure_ascii=False,
                indent=2,
                allow_nan=False,
            ),
        )
        z.writestr(
            "README.txt",
            (
                "case.json 可重新导入工作台复算。evaluation.json 保留完整计算快"
                "照。原始资料按 SHA-256 对照项目资料库；本包不重复包含原件。\n"
            ),
        )
    return stream.getvalue()
