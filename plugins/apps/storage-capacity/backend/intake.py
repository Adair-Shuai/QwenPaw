# -*- coding: utf-8 -*-
"""Read evidence, preserve locations, and map only user-selected table \
columns."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from pathlib import Path
import posixpath
from zipfile import ZipFile, BadZipFile
from xml.etree import ElementTree as ET

from .catalog import FIELDS, new_unit, normalize_case

MAX_FILE = 50 * 1024 * 1024
MAX_CELLS = 200000
NAMESPACES = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def slide_names(archive):
    """Slide order follows presentation relationships, not slide filenames."""
    if "ppt/presentation.xml" not in archive.namelist():
        return sorted(
            [
                n
                for n in archive.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)
            ],
            key=lambda n: int(re.search(r"slide(\d+)", n)[1]),
        )
    relations = ET.fromstring(archive.read("ppt/_rels/presentation.xml.rels"))
    targets = {
        r.attrib["Id"]: r.attrib["Target"]
        for r in relations
        if r.attrib.get("TargetMode") != "External"
    }
    presentation = ET.fromstring(archive.read("ppt/presentation.xml"))
    names = []
    for node in presentation.findall(
        ".//{http://schemas.openxmlformats.org/presentationml/2006/main}sldId",
    ):
        target = targets[
            node.attrib[
                (
                    "{http://schemas.openxmlformats.org/officeDoc"
                    "ument/2006/relationships}id"
                )
            ]
        ]
        name = (
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.normpath(posixpath.join("ppt", target))
        )
        if not name.startswith("ppt/slides/"):
            raise ValueError("幻灯片关系无效")
        names.append(name)
    return names


def numeric(value):
    if value in (None, "", "/", "—", "--", "未测试"):
        return None
    if isinstance(value, bool):
        raise ValueError("布尔值不是工程数值")
    text = str(value).strip().replace(",", "")
    if not re.fullmatch(
        r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?",
        text,
    ):
        raise ValueError(f"无法确认为数值：{str(value)[:60]}")
    number = float(text)
    if not __import__("math").isfinite(number):
        raise ValueError("数值必须有限")
    return number


def _zip_guard(raw):
    with ZipFile(io.BytesIO(raw)) as archive:
        infos = archive.infolist()
        if (
            len(infos) > 20000
            or sum(i.file_size for i in infos) > 250 * 1024 * 1024
        ):
            raise ValueError("压缩文件展开后过大，请拆分资料")


def _extract_spreadsheet(
    raw,
    result,
    table,
):
    """Extract spreadsheet."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        for ws in wb:
            if (ws.max_row or 0) * (ws.max_column or 0) > MAX_CELLS:
                raise ValueError("工作表范围过大，请清理空白格式或拆分")
            rows = [
                [v.isoformat() if hasattr(v, "isoformat") else v for v in row]
                for row in ws.iter_rows(values_only=True)
            ]
            table(ws.title, rows)
    finally:
        wb.close()
    result["warnings"].append(
        "公式读取 Excel 保存的缓存值；空缓存保持缺失。合并表头需选择实际表头行，缺测符号不按零处理。",
    )


def _extract_delimited_text(
    suffix,
    filename,
    raw,
    table,
):
    """Extract delimited text."""
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("gb18030")
    table(
        filename,
        list(
            csv.reader(
                io.StringIO(text),
                delimiter="\t" if suffix == ".tsv" else ",",
            ),
        ),
    )


def _extract_presentation(
    raw,
    result,
    table,
):
    """Extract presentation."""
    with ZipFile(io.BytesIO(raw)) as archive:
        names = slide_names(archive)
        for i, name in enumerate(names, 1):
            tree = ET.fromstring(archive.read(name))
            text = "\n".join(
                "".join(p.itertext())
                for p in tree.findall(".//a:t", NAMESPACES)
            )
            result["pages"].append(
                {"label": f"第 {i} 页", "text": text[:30000]},
            )
            for j, t in enumerate(
                tree.findall(".//a:tbl", NAMESPACES),
                1,
            ):
                rows = [
                    [
                        "".join(c.itertext()).strip()
                        for c in row.findall("a:tc", NAMESPACES)
                    ]
                    for row in t.findall("a:tr", NAMESPACES)
                ]
                # itertext above includes formatting whitespace but
                # no hidden relationships.
                table(f"第 {i} 页 表 {j}", rows)
    result["warnings"].append("已提取可编辑文字与表格；图像内的公式、曲线和扫描内容需对照原件核实。")


def _extract_pdf(
    raw,
    result,
):
    """Extract pdf."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    if len(reader.pages) > 500:
        raise ValueError("PDF 超过 500 页，请拆分")
    for i, page in enumerate(reader.pages, 1):
        result["pages"].append(
            {
                "label": f"第 {i} 页",
                "text": (page.extract_text() or "")[:30000],
            },
        )
    result["warnings"].append(
        "PDF 文字不保证保留表格列对应关系；请核对原页后填入，扫描页需 OCR 或手工转录。",
    )


def _extract_word_document(
    raw,
    result,
    table,
):
    """Extract word document."""
    from docx import Document

    doc = Document(io.BytesIO(raw))
    result["pages"].append(
        {
            "label": "正文（无物理页码）",
            "text": "\n".join(p.text for p in doc.paragraphs)[:100000],
        },
    )
    for i, t in enumerate(doc.tables, 1):
        table(f"表 {i}", [[c.text for c in r.cells] for r in t.rows])


def _extract_json_document(
    raw,
    result,
    filename,
    table,
):
    """Extract json document."""
    parsed = json.loads(raw.decode("utf-8-sig"))
    candidate = (
        parsed.get("case", parsed) if isinstance(parsed, dict) else None
    )
    if isinstance(candidate, dict) and candidate.get("schema_version") == 2:
        result["case_candidate"] = normalize_case(candidate)
    elif (
        isinstance(parsed, list)
        and parsed
        and all(isinstance(r, dict) for r in parsed)
    ):
        keys = list(dict.fromkeys(k for row in parsed for k in row))
        table(
            filename,
            [keys] + [[r.get(k) for k in keys] for r in parsed],
        )
    else:
        result["pages"].append(
            {
                "label": "JSON",
                "text": json.dumps(
                    parsed,
                    ensure_ascii=False,
                    indent=2,
                )[:100000],
            },
        )


def _table_collector(result):
    """Own the cumulative source-table cell budget for one document."""
    cell_count = 0

    def table(label, rows, locations=None):
        nonlocal cell_count
        cell_count += sum(len(r) for r in rows)
        if cell_count > MAX_CELLS or len(rows) > 10000:
            raise ValueError("表格超过 20 万单元格或 1 万行，请按周期和用途拆分")
        result["tables"].append(
            {
                "id": str(len(result["tables"])),
                "label": label,
                "rows": rows,
                "locations": locations
                or [f"{label} 第 {i + 1} 行" for i in range(len(rows))],
            },
        )

    return table


def _mark_empty_document(result):
    """Require manual review for documents without extracted evidence."""
    if (
        not result["tables"]
        and not any(p["text"].strip() for p in result["pages"])
        and not result.get("case_candidate")
    ):
        result["status"] = "manual_review"


def extract(filename, raw):
    if len(raw) > MAX_FILE:
        raise ValueError("单份资料不能超过 50 MB")
    filename = Path(filename.replace("\\", "/")).name
    suffix = Path(filename).suffix.lower()
    digest = hashlib.sha256(raw).hexdigest()
    result = {
        "id": digest[:24],
        "sha256": digest,
        "name": filename,
        "size_bytes": len(raw),
        "tables": [],
        "pages": [],
        "warnings": [],
        "status": "extracted",
    }
    table = _table_collector(result)

    try:
        if suffix in {".xlsx", ".docx", ".pptx"}:
            _zip_guard(raw)
        if suffix == ".xlsx":
            _extract_spreadsheet(
                raw,
                result,
                table,
            )
        elif suffix in {".csv", ".tsv"}:
            _extract_delimited_text(
                suffix,
                filename,
                raw,
                table,
            )
        elif suffix == ".pptx":
            _extract_presentation(
                raw,
                result,
                table,
            )
        elif suffix == ".pdf":
            _extract_pdf(
                raw,
                result,
            )
        elif suffix == ".docx":
            _extract_word_document(
                raw,
                result,
                table,
            )
        elif suffix == ".json":
            _extract_json_document(
                raw,
                result,
                filename,
                table,
            )
        elif suffix in {".txt", ".md"}:
            result["pages"].append(
                {"label": "文本", "text": raw.decode("utf-8-sig")[:100000]},
            )
        else:
            result["status"] = "manual_review"
            result["warnings"].append(
                "该格式仅保存原件；请提供 XLSX、CSV、JSON、PPTX、PDF 或 DOCX 以读取内容。",
            )
    except (ImportError, BadZipFile, UnicodeDecodeError) as exc:
        result["status"] = "manual_review"
        result["warnings"].append(f"未能自动读取，请核对原件：{type(exc).__name__}")
    _mark_empty_document(result)
    return result


MAPPING_FIELDS = {
    "units": ["name"] + list(FIELDS),
    "series": ["cumulative_produced", "pressure", "z", "phase"],
    "pvt": ["pressure", "z"],
    "wells": ["name", "a", "b", "reservoir_pressure", "bottomhole_pressure"],
    "schedule": ["day", "rate"],
    "history": [
        "cycle_id",
        "book_inventory",
        "effective_inventory",
        "geological_swing",
        "deliverable_swing",
        "peak_rate",
    ],
    "history_match": ["observed", "simulated"],
}


def _append_mapped_rows(
    table,
    start,
    end,
    document,
    mapping,
    factors,
    allowed_factors,
    kind,
    records,
    errors,
):
    """Append mapped rows."""
    for index, row in enumerate(table["rows"][start:end], start):
        if not any(v not in (None, "") for v in row):
            continue
        source = f"{document['name']} · {table['locations'][index]}"
        record = {"source": source, "_evidence": {}}
        for field, col in mapping.items():
            if isinstance(col, bool) or not isinstance(col, int) or col < 0:
                raise ValueError("列序号必须为非负整数")
            raw = row[col] if col < len(row) else None
            try:
                value = (
                    str(raw).strip()
                    if field in {"name", "cycle_id", "phase"}
                    and raw is not None
                    else numeric(raw)
                )
                factor = float(factors.get(field, 1))
                if factor not in allowed_factors:
                    raise ValueError("不支持的换算倍数")
                if isinstance(value, (int, float)):
                    value *= factor
                record[field] = value
                record["_evidence"][field] = {
                    "document_id": document["id"],
                    "sha256": document["sha256"],
                    "location": f"{table['label']}!R{index + 1}C{col + 1}",
                    "raw": raw,
                    "factor": factor,
                }
            except (ValueError, TypeError) as exc:
                errors.append(f"{source} 第 {col + 1} 列：{exc}")
        if kind == "units":
            if not record.get("name"):
                errors.append(f"{source} 缺少单元名称")
            unit = new_unit(str(record.get("name") or "未命名单元"))
            unit.update(
                source=source,
                parameters={k: v for k, v in record.items() if k in FIELDS},
                evidence=record["_evidence"],
            )
            records.append(unit)
        else:
            records.append(record)


def map_table(document, body):
    kind = body.get("kind")
    if kind not in MAPPING_FIELDS:
        raise ValueError("请选择受支持的资料类型")
    table = next(
        (
            t
            for t in document["tables"]
            if t["id"] == str(body.get("table_id"))
        ),
        None,
    )
    if table is None:
        raise ValueError("表格不存在")
    mapping = body.get("mapping") or {}
    if not mapping or any(k not in MAPPING_FIELDS[kind] for k in mapping):
        raise ValueError("字段映射为空或含未知字段")
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("同一列不能同时映射到多个参数")
    start = int(body.get("start_row", 1))
    end = int(body.get("end_row", len(table["rows"])))
    if start < 0 or end > len(table["rows"]) or start >= end:
        raise ValueError("数据行范围无效")
    factors = body.get("factors") or {}
    allowed_factors = {1.0, 0.001, 0.01, 0.0001, 1e-8, 1e4, 1e8, 1000.0}
    records, errors = [], []
    _append_mapped_rows(
        table,
        start,
        end,
        document,
        mapping,
        factors,
        allowed_factors,
        kind,
        records,
        errors,
    )
    if errors:
        return {"records": [], "errors": errors, "kind": kind}
    return {"records": records, "errors": [], "kind": kind}
