# -*- coding: utf-8 -*-
"""Excel-backed, deterministic document forms. No model service is involved."""

from __future__ import annotations

import io
from decimal import Decimal, InvalidOperation
import re
from datetime import date, datetime
from zipfile import ZipFile, BadZipFile

from docxtpl import DocxTemplate
from jinja2 import StrictUndefined, nodes
from jinja2.sandbox import SandboxedEnvironment
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from lxml import etree

MAX_FILE = 15 * 1024 * 1024
PATH = re.compile(r"^[^\W\d]\w*(?:\.[^\W\d]\w*)*$", re.UNICODE)
INVITATION_KEYS = {
    "invitation.recipient",
    "events.invitation_time",
    "invitation.start",
    "invitation.end",
}
PARTICIPANT_KEYS = {"participants", "participant_groups"}


def template_metadata(name, index=1):
    """Separate formal document names from source-file working notes."""
    titles = (
        ("外协必要性说明", ("外协必要性说明",)),
        ("科研计划填报", ("科研计划填报",)),
        ("论证意见及签字表", ("论证意见",)),
        ("非招标方案填报", ("非招标方案填报",)),
        ("商务谈判文件", ("商务谈判文件", "谈判文件")),
        ("谈判邀请函", ("谈判邀请函", "邀请函")),
        ("投标报价一览表", ("投标报价一览", "谈判报价一览")),
        ("综合评分表", ("综合评分打分表", "综合评分表")),
        ("综合评分汇总表", ("综合评分总表", "综合评分汇总表")),
    )
    extension = ".xlsx" if name.lower().endswith(".xlsx") else ".docx"
    title = next(
        (
            title
            for title, aliases in titles
            if any(a in name for a in aliases)
        ),
        None,
    )
    if title is None:
        title = re.sub(r"\.(?:docx|xlsx)$", "", name, flags=re.I)
        title = re.sub(r"^\d+[\s_.、-]+", "", title)
        title = re.sub(r"[ _-]*(?:示例)?模板$", "", title).strip() or "文书"
    return {
        "title": title,
        "name": f"{index:02d}_{title}_模板{extension}",
        "extension": extension,
        "print_copies": 7 if title == "综合评分表" else 1,
        "print_note": (
            "每位专家一份，得分及签字现场填写" if title == "综合评分表" else ""
        ),
    }


GROUPS = {
    "project": "项目资料",
    "organization": "单位与联系",
    "org": "单位与联系",
    "contact": "单位与联系",
    "purchaser": "单位与联系",
    "dates": "日期安排",
    "events": "日期安排",
    "qualification": "资格要求",
    "submission": "日期安排",
    "research": "研究内容",
    "deliverables": "成果与考核",
    "metrics": "成果与考核",
    "workload": "研究内容",
    "milestones": "进度安排",
    "place": "地点安排",
    "budget": "费用信息",
}


def safe_key(value):
    value = str(value or "").strip()
    if not PATH.fullmatch(value) or any(
        p.startswith("_") for p in value.split(".")
    ):
        raise ValueError(
            f"变量名称不合规：{value}。请使用名称或以点分隔的名称。",
        )
    return value


def inspect_zip(content: bytes, required: str):
    if len(content) > MAX_FILE:
        raise ValueError("单个文件不能超过 15 MB")
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            items = archive.infolist()
            if (
                len(items) > 3000
                or sum(i.file_size for i in items) > 60 * 1024 * 1024
            ):
                raise ValueError("文件解压后过大，请拆分后导入")
            if required not in archive.namelist():
                raise ValueError("文件内容与扩展名不一致")
            if any("vbaproject" in i.filename.lower() for i in items):
                raise ValueError("请使用不含宏的模板")
    except BadZipFile as exc:
        raise ValueError("文件损坏或不是有效的 Office 文件") from exc


def text(value):
    if isinstance(value, datetime):
        return (
            value.strftime("%Y-%m-%d %H:%M")
            if value.time().isoformat() != "00:00:00"
            else value.strftime("%Y-%m-%d")
        )
    if isinstance(value, date):
        return value.isoformat()
    return "" if value is None else str(value)


def kind_for(key, label, value):
    if "year" in key or "年份" in label:
        return "text"
    if any(
        x in key + label
        for x in (
            "日期",
            "时间",
            "deadline",
            "_time",
            "dates.",
            ".start",
            ".end",
        )
    ):
        # Chinese date phrases and ranges are retained as editable text.
        if not value or re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            return "date"
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(?::\d{2})?", value):
            return "datetime-local"
        if parsed_date(value):
            return "datetime-local" if re.search(r"[:时]", value) else "date"
    return "text"


def recipient_defaults(recipients):
    """Materialize linked times; never infer a link from an explicit value."""
    result = [dict(recipient) for recipient in recipients]
    if not result:
        return result
    for index, recipient in enumerate(result):
        for key in ("start", "end"):
            flag = "follow_" + key
            follow = recipient.get(flag)
            if follow is None:
                follow = not bool(recipient.get(key, "").strip())
            recipient[flag] = bool(index and follow)
            if recipient[flag]:
                recipient[key] = result[0].get(key, "")
    return result


def read_excel(content: bytes):
    # Several independent workbook sections are validated in one pass.
    # pylint: disable=too-many-branches,too-many-statements
    inspect_zip(content, "xl/workbook.xml")
    book = load_workbook(io.BytesIO(content), data_only=False, read_only=True)
    try:

        def rows_for(sheet, max_rows, max_columns):
            # Some valid writers omit dimensions; others write stale bounds.
            # Scan actual rows with hard limits instead of trusting metadata.
            sheet.reset_dimensions()
            rows = []
            for index, row in enumerate(sheet.iter_rows(), 1):
                if index > max_rows or len(row) > max_columns:
                    raise ValueError(
                        f"{sheet.title}最多支持 {max_rows - 1} 行、{max_columns} 列",
                    )
                rows.append(row)
            if not rows:
                raise ValueError(f"{sheet.title}为空，请添加表头和数据")
            return rows

        if "变量配置" not in book.sheetnames:
            raise ValueError(
                "Excel 缺少“变量配置”工作表，首行为：变量名称、中文名称、值",
            )
        sheet = book["变量配置"]
        rows = rows_for(sheet, 1001, 30)
        headers = [text(c.value).strip() for c in rows[0]]
        if len(set(headers) & {"变量名称", "中文名称", "值"}) != 3:
            raise ValueError("变量配置首行必须包含：变量名称、中文名称、值")

        def cell(row, name, default=""):
            if name not in headers:
                return default
            index = headers.index(name)
            if index >= len(row):
                return ""
            c = row[index]
            if c.data_type == "f":
                raise ValueError(
                    f"{sheet.title}!{c.coordinate} 使用了公式，请改为明确的值",
                )
            return text(c.value).strip()

        fields = []
        seen = set()
        for row in rows[1:]:
            if not any(c.value is not None for c in row):
                continue
            key = safe_key(cell(row, "变量名称"))
            if key.split(".", maxsplit=1)[0] in PARTICIPANT_KEYS:
                raise ValueError(
                    "participants 和 participant_groups 为单位名单保留变量，"
                    "请勿在 Excel 中配置",
                )
            if key in seen or any(
                k.startswith(key + ".") or key.startswith(k + ".")
                for k in seen
            ):
                raise ValueError(f"变量重复或层级冲突：{key}")
            seen.add(key)
            label, value = cell(row, "中文名称") or key, cell(row, "值")
            kind = cell(row, "类型") or kind_for(key, label, value)
            if kind not in {
                "text",
                "textarea",
                "date",
                "datetime-local",
                "number",
            }:
                raise ValueError(f"{label} 的类型不支持：{kind}")
            if (
                value
                and kind in {"date", "datetime-local"}
                and parsed_date(value)
            ):
                value = parsed_date(value).strftime(
                    (
                        "%Y-%m-%dT%H:%M"
                        if kind == "datetime-local"
                        else "%Y-%m-%d"
                    ),
                )
            fields.append(
                {
                    "key": key,
                    "label": label,
                    "value": value,
                    "group": cell(row, "分组")
                    or GROUPS.get(key.split(".", 1)[0], "其他资料"),
                    "kind": kind,
                    "required": cell(row, "必填", "是")
                    not in {"否", "false", "0"},
                },
            )
        if not fields:
            raise ValueError("变量配置中没有可填写的变量")
        recipients = []
        if "受邀单位" in book.sheetnames:
            sheet = book["受邀单位"]
            data = rows_for(sheet, 201, 20)
            # Parenthesized variable names and format hints are descriptions.
            headers = [
                re.split(r"[（(]", text(c.value), maxsplit=1)[0].strip()
                for c in data[0]
            ]
            aliases = {
                "name": ["单位名称", "受邀单位", "受邀单位名称", "名称"],
                "start": [
                    "领取开始时间",
                    "邀请函领取开始时间",
                    "邀请函领取开始",
                    "领取起始时间",
                    "领取起始",
                    "开始时间",
                    "领取开始",
                ],
                "end": [
                    "领取结束时间",
                    "邀请函领取结束时间",
                    "邀请函领取结束",
                    "领取截止时间",
                    "领取截止",
                    "结束时间",
                    "领取结束",
                ],
            }
            if not any(n in headers for n in aliases["name"]):
                raise ValueError("受邀单位工作表缺少“单位名称”列")
            for row_number, row in enumerate(data[1:], 2):
                if not any(c.value is not None for c in row):
                    continue
                record = {}
                for key, names in aliases.items():
                    record[key] = next(
                        (cell(row, n) for n in names if n in headers),
                        "",
                    )
                    if key != "name" and record[key]:
                        parsed = parsed_date(record[key])
                        if parsed is None:
                            raise ValueError(
                                f"受邀单位第 {row_number} 行的领取时间无法识别，请填写明确的日期时间",
                            )
                        record[key] = parsed.strftime("%Y-%m-%dT%H:%M")
                # Keep numbered empty slots for later completion, but not a
                # trailing note stored outside the mapped input columns.
                if not (
                    any(record.values())
                    or ("序号" in headers and cell(row, "序号").isdigit())
                    or any(
                        cell(row, column)
                        for column in ("开始时间跟随首家", "结束时间跟随首家")
                    )
                ):
                    continue
                for key, column in [
                    ("start", "开始时间跟随首家"),
                    ("end", "结束时间跟随首家"),
                ]:
                    flag = cell(row, column).lower()
                    if flag:
                        if flag not in {
                            "是",
                            "否",
                            "true",
                            "false",
                            "1",
                            "0",
                        }:
                            raise ValueError(
                                f"受邀单位第 {row_number} 行的{column}请填写“是”或“否”",
                            )
                        record["follow_" + key] = flag in {
                            "是",
                            "true",
                            "1",
                        }
                recipients.append(record)
        rules = []
        if "校核规则" in book.sheetnames:
            sheet = book["校核规则"]
            data = rows_for(sheet, 201, 30)
            headers = [text(c.value).strip() for c in data[0]]
            if not {"校核项", "开始变量", "结束变量"}.issubset(headers):
                raise ValueError(
                    "校核规则首行需包含：校核项、开始变量、结束变量",
                )
            for row in data[1:]:
                if any(c.value is not None for c in row):
                    rules.append(
                        {
                            "label": cell(row, "校核项"),
                            "start": safe_key(cell(row, "开始变量")),
                            "end": safe_key(cell(row, "结束变量")),
                            "equal": cell(row, "允许相同", "是") != "否",
                        },
                    )
        # 日期校核 may contain Excel formulas; the application always recalculates
        # from editable source values instead of trusting stale cached results.
        return {
            "fields": fields,
            "recipients": recipient_defaults(recipients),
            "rules": rules,
        }
    finally:
        book.close()


class FormEnvironment(SandboxedEnvironment):
    def getattr(self, obj, attribute):
        # Excel variables such as project.items are data, not dict methods.
        if isinstance(obj, dict) and attribute in obj:
            return obj[attribute]
        return super().getattr(obj, attribute)


def environment():
    env = FormEnvironment(undefined=StrictUndefined, autoescape=True)
    env.globals.clear()
    return env


def _visit_template(node, undeclared, paths):
    if isinstance(node, nodes.Getattr):
        parts = []
        cursor = node
        while isinstance(cursor, nodes.Getattr):
            parts.insert(0, cursor.attr)
            cursor = cursor.node
        if isinstance(cursor, nodes.Name) and cursor.name in undeclared:
            paths.add(safe_key(".".join([cursor.name] + parts)))
            return
    if isinstance(node, nodes.Name) and node.name in undeclared:
        paths.add(safe_key(node.name))
    if isinstance(node, (nodes.Call, nodes.Getitem)):
        raise ValueError(
            "模板请使用名称或点分隔的变量，不支持函数调用和下标取值",
        )
    for child in node.iter_child_nodes():
        _visit_template(child, undeclared, paths)


def template_info(content: bytes, name: str, global_keys=()):
    if name.lower().endswith(".xlsx"):
        # Plugins load as dynamic packages; static lint cannot resolve them.
        # pylint: disable=import-outside-toplevel,relative-beyond-top-level
        from .spreadsheets import template_keys

        inspect_zip(content, "xl/workbook.xml")
        return {
            **template_metadata(name),
            "keys": sorted(template_keys(content, safe_key)),
            "per_recipient": False,
        }
    inspect_zip(content, "word/document.xml")
    template = DocxTemplate(io.BytesIO(content))
    template.init_docx()
    paths = set()
    env = environment()
    with ZipFile(io.BytesIO(content)) as archive:
        for member in archive.namelist():
            if not re.fullmatch(
                r"word/(document|header\d*|footer\d*|footnotes|endnotes)\.xml",
                member,
            ):
                continue
            xml = template.patch_xml(archive.read(member).decode("utf-8"))
            tree = env.parse(xml)
            from jinja2 import meta

            undeclared = meta.find_undeclared_variables(tree)
            _visit_template(tree, undeclared, paths)
    if not paths:
        raise ValueError(
            f"{name} 没有找到变量，请在 Word 中使用 {{{{ project.name }}}} 这样的占位符",
        )
    return {
        **template_metadata(name),
        "keys": sorted(paths),
        "per_recipient": bool(
            paths & (INVITATION_KEYS - {"events.invitation_time"}),
        )
        or (
            "events.invitation_time" in paths
            and "events.invitation_time" not in global_keys
        ),
    }


def parsed_date(value):
    period = re.search(r"凌晨|早上|上午|中午|下午|晚上|夜间|午夜", str(value))
    value = (
        str(value)
        .strip()
        .replace("：", ":")
        .replace("年", "-")
        .replace("月", "-")
        .replace("日", " ")
        .replace("时", ":")
        .replace("分", "")
        .strip()
    )
    if period:
        value = value.replace(period[0], "")
    value = re.sub(r"\s+", " ", value)
    match = re.fullmatch(
        r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})"
        r"(?:[ T](\d{1,2}):(\d{1,2})(?::(\d{1,2}))?)?",
        value,
    )
    if match:
        try:
            parts = [int(v or 0) for v in match.groups()]
            if period:
                if match[4] is None:
                    return None
                hour = parts[3]
                if period[0] in {"下午", "晚上", "夜间"} and 1 <= hour < 12:
                    parts[3] += 12
                elif period[0] in {"凌晨", "午夜"} and hour == 12:
                    parts[3] = 0
                elif period[0] in {"早上", "上午"} and hour >= 12:
                    return None
            return datetime(*parts)
        except ValueError:
            return None
    try:
        result = datetime.fromisoformat(value)
        return result if result.tzinfo is None else None
    except ValueError:
        return None


def parsed_interval(value):
    """Return explicit endpoints while keeping the original document text."""
    parts = re.split(r"\s*(?:~|～|至|到|—|–)\s*", str(value).strip())
    if len(parts) == 1:
        point = parsed_date(parts[0])
        return (point, point) if point else None
    if len(parts) == 2:
        start, end = (parsed_date(part) for part in parts)
        if (
            start
            and not end
            and re.fullmatch(
                r"(?:凌晨|早上|上午|中午|下午|晚上|夜间|午夜)?\s*"
                r"\d{1,2}[:：时]\d{1,2}(?:分|[:：]\d{1,2})?",
                parts[1],
            )
        ):
            end_text = parts[1]
            period = re.search(r"下午|晚上|夜间", parts[0])
            if period and re.match(r"\d{1,2}[:：时]", end_text):
                if int(re.match(r"\d{1,2}", end_text)[0]) < 12:
                    end_text = period[0] + end_text
            end = parsed_date(start.strftime("%Y年%m月%d日") + end_text)
        if start and end:
            return start, end
    return None


def day_period(hour):
    if hour < 6:
        return "凌晨"
    if hour < 12:
        return "上午"
    if hour == 12:
        return "中午"
    return "下午" if hour < 18 else "晚上"


def document_date(value):
    """Format document text without rewriting source date-control values."""
    interval = parsed_interval(value)
    if not interval:
        return str(value or "")
    start, end = interval
    has_clock = bool(re.search(r"\d[:：时]\d", str(value)))

    def date_text(point):
        return f"{point.year}年{point.month}月{point.day}日"

    def clock(point):
        return point.strftime("%H:%M:%S" if point.second else "%H:%M")

    first = date_text(start)
    if has_clock:
        first += day_period(start.hour) + clock(start)
    if start == end:
        return first
    if start.date() == end.date() and has_clock:
        return (
            first
            + "至"
            + (
                ""
                if day_period(start.hour) == day_period(end.hour)
                else day_period(end.hour)
            )
            + clock(end)
        )
    return (
        first
        + "至"
        + date_text(end)
        + (day_period(end.hour) + clock(end) if has_clock else "")
    )


def document_field_value(field):
    value = field["value"]
    if (
        field.get("kind") in {"date", "datetime-local"}
        or re.search(
            r"(^dates\.|deadline$|_time$|\.start$|\.end$)",
            field["key"],
        )
        or any(
            word in field.get("label", "") for word in ("日期", "时间", "截止")
        )
    ):
        return document_date(value)
    if field["key"].startswith("milestones.") or "进度安排" in field.get(
        "label",
        "",
    ):
        value = re.sub(
            r"(?<!\d)\d{4}[-/.]\d{1,2}[-/.]\d{1,2}(?!\d)",
            lambda match: document_date(match[0]),
            value,
        )
        value = re.sub(
            r"(?<!\d)(\d{4})[-/.](0?[1-9]|1[0-2])(?!\d|[-/.]\d{1,2}(?!\d))",
            lambda match: f"{int(match[1])}年{int(match[2])}月",
            value,
        )
        value = re.sub(r"月\s*[-~～—–]\s*(?=\d{4}年)", "月至", value)
    return value


def date_roles(fields):
    aliases = {
        "issue": ["dates.issue", "dates.invitation_issue"],
        "qa": [
            "events.qa_deadline",
            "events.qa_time",
            "dates.qa_deadline",
            "events.questions_time",
        ],
        "response": [
            "events.response_deadline",
            "dates.response_deadline",
            "submission.deadline",
        ],
        "negotiation": ["events.negotiation_time", "dates.negotiation"],
    }
    labels = {
        "issue": ["落款日期", "发文日期", "文件日期", "发文时间"],
        "qa": ["答疑", "澄清截止"],
        "response": [
            "响应截止",
            "响应文件截止",
            "响应文件递交截止",
            "响应文件提交截止",
        ],
        "negotiation": ["谈判时间", "商务谈判日期", "谈判日期"],
    }
    result = {}
    for role, keys in aliases.items():
        field = next((f for f in fields if f["key"] in keys), None) or next(
            (f for f in fields if any(x in f["label"] for x in labels[role])),
            None,
        )
        if field:
            result[role] = field["key"]
    return result


def checks(case):
    # Required fields, templates, shared dates and budgets each add rules.
    # pylint: disable=too-many-branches,too-many-statements
    fields, recipients = case["fields"], case["recipients"]
    values = {f["key"]: f["value"] for f in fields}
    field_labels = {f["key"]: f["label"] for f in fields}
    results = []

    def add(label, status, detail, field=""):
        results.append(
            {
                "label": label,
                "status": status,
                "detail": detail,
                "field": field,
            },
        )

    for f in fields:
        if f.get("options") and f["value"] and f["value"] not in f["options"]:
            add(f["label"], "error", "请从下拉选项中选择", f["key"])
        if (
            f["key"]
            in {
                "plan.budget_wan",
                "selection.budget_wan",
                "funding.amount_wan",
            }
            and f["value"]
        ):
            try:
                amount = Decimal(f["value"])
                valid = (
                    amount.is_finite()
                    and amount >= 0
                    and len(f["value"]) <= 50
                )
            except InvalidOperation:
                valid = False
            if not valid:
                add(f["label"], "error", "请按万元填写非负数值", f["key"])
        if f["key"] == "selection.overview" and len(f["value"].strip()) > 500:
            add(f["label"], "error", "系统填报要求不超过 500 字", f["key"])
        if (
            f["key"] == "selection.year"
            and f["value"]
            and not re.fullmatch(r"[12]\d{3}", f["value"])
        ):
            add(f["label"], "error", "请填写四位计划年度", f["key"])
        if f["required"] and not str(f["value"]).strip():
            add(f["label"], "missing", "尚未填写", f["key"])
        elif (
            f["value"]
            and f["kind"] in {"date", "datetime-local"}
            and not parsed_date(f["value"])
        ):
            add(f["label"], "error", "日期格式无法识别", f["key"])
    for key in {k for t in case["templates"] for k in t["keys"]}:
        if (
            key not in INVITATION_KEYS | PARTICIPANT_KEYS
            and key not in values
            and not any(k.startswith(key + ".") for k in values)
        ):
            add(key, "missing", "文书模板使用了此变量，Excel 中没有配置", key)
    if not case["templates"]:
        add("文书模板", "missing", "请至少导入一个 Word 或 Excel 模板")
    count = sum(
        len(recipients) if t["per_recipient"] else 1 for t in case["templates"]
    )
    if count > 200:
        add(
            "生成数量",
            "error",
            "每批最多生成 200 份文书，请减少受邀单位或模板数量",
        )
    per_recipient = any(t["per_recipient"] for t in case["templates"])
    needs_participants = any(
        PARTICIPANT_KEYS & set(t["keys"]) for t in case["templates"]
    )
    if (per_recipient or needs_participants) and not recipients:
        add(
            "受邀单位", "missing", "按单位生成或汇总的文书需要至少一家受邀单位"
        )
    roles = date_roles(fields)
    rules = list(case["rules"])
    chain = [
        r for r in ("issue", "qa", "response", "negotiation") if r in roles
    ]
    for a, b in zip(chain, chain[1:]):
        operator = "<" if a == "response" else "≤"
        rules.append(
            {
                "label": (
                    f"{field_labels[roles[a]]} {operator} "
                    f"{field_labels[roles[b]]}"
                ),
                "start": roles[a],
                "end": roles[b],
                "equal": a != "response",
            },
        )

    def compare(label, a, b, equal=True, field=""):
        left, right = parsed_interval(a), parsed_interval(b)
        if not a or not b:
            add(label, "missing", "日期尚未补齐", field)
        elif not left or not right:
            add(
                label,
                "error",
                "请填写明确日期或完整时间段，如 2026年7月25日下午14:30至17:30",
                field,
            )
        elif left[1] <= right[0] if equal else left[1] < right[0]:
            add(label, "ok", f"{document_date(a)} → {document_date(b)}", field)
        else:
            add(
                label,
                "error",
                f"时间顺序冲突：{document_date(a)} → {document_date(b)}",
                field,
            )

    check_range_order(values, field_labels, roles, add)
    overall = values.get("events.invitation_time")
    overall_interval = parsed_interval(overall)
    if overall is not None:
        if "issue" in roles:
            compare(
                "落款日期不晚于总体发放开始",
                values[roles["issue"]],
                overall,
                field="events.invitation_time",
            )
        for role in ("qa", "response", "negotiation"):
            if role in roles:
                compare(
                    f"总体发放结束至{field_labels[roles[role]]}",
                    overall,
                    values[roles[role]],
                    field=roles[role],
                )

    for rule in rules:
        compare(
            rule["label"],
            values.get(rule["start"]),
            values.get(rule["end"]),
            rule["equal"],
            rule["end"],
        )
    if (
        "dates.invitation_issue" in values
        and "issue" in roles
        and roles["issue"] != "dates.invitation_issue"
    ):
        add(
            "邀请函与谈判文件落款一致",
            (
                "ok"
                if parsed_date(values["dates.invitation_issue"])
                == parsed_date(values[roles["issue"]])
                and parsed_date(values["dates.invitation_issue"]) is not None
                else "error"
            ),
            "两份文书应使用同一落款日期",
            "dates.invitation_issue",
        )
    if per_recipient or recipients:
        check_recipients(case, roles, add, compare, overall_interval)
    return results


def check_range_order(values, labels, roles, add):
    # Verify range order before comparing its end with the next event's start.
    range_keys = set(roles.values()) | (
        {"events.invitation_time"} & values.keys()
    )
    for key in sorted(range_keys):
        value = values[key]
        interval = parsed_interval(value)
        if value and not interval:
            add(
                labels[key],
                "error",
                "日期或时间段格式无法识别，请填写完整起止日期时间",
                key,
            )
        elif interval and interval[0] != interval[1]:
            add(
                labels[key] + " · 起止顺序",
                "ok" if interval[0] <= interval[1] else "error",
                document_date(value),
                key,
            )


def check_recipients(case, roles, add, compare, overall_interval):
    values = {f["key"]: f["value"] for f in case["fields"]}
    field_labels = {f["key"]: f["label"] for f in case["fields"]}
    seen = set()
    for i, r in enumerate(case["recipients"]):
        title = r["name"].strip() or f"第 {i + 1} 家单位"
        if not r["name"].strip():
            add(title, "missing", "请填写单位名称", f"recipient-{i}")
        elif title in seen:
            add(title, "error", "单位名称重复，请核对", f"recipient-{i}")
        seen.add(title)
        compare(
            f"{title} · 领取起止",
            r["start"],
            r["end"],
            field=f"recipient-{i}",
        )
        if "issue" in roles:
            compare(
                f"{title} · 发文至领取",
                values[roles["issue"]],
                r["start"],
                field=f"recipient-{i}",
            )
        for role in ("qa", "response", "negotiation"):
            if role in roles:
                endpoint = parsed_interval(values[roles[role]])
                compare(
                    f"{title} · 领取结束至{field_labels[roles[role]]}",
                    r["end"],
                    (
                        endpoint[0].isoformat()
                        if endpoint
                        else values[roles[role]]
                    ),
                    field=f"recipient-{i}",
                )
        if overall_interval:
            compare(
                f"{title} · 领取开始处于总体发放期",
                overall_interval[0].isoformat(),
                r["start"],
                field=f"recipient-{i}",
            )
            compare(
                f"{title} · 领取结束处于总体发放期",
                r["end"],
                overall_interval[1].isoformat(),
                field=f"recipient-{i}",
            )


def render_template(
    content,
    fields,
    recipient=None,
    participants=(),
    extension=".docx",
):
    if extension == ".xlsx":
        # Plugins load as dynamic packages; static lint cannot resolve them.
        # pylint: disable=import-outside-toplevel,relative-beyond-top-level
        from .spreadsheets import fill

        inspect_zip(content, "xl/workbook.xml")
        return fill(content, fields)
    context = {}
    for f in fields:
        cursor = context
        parts = f["key"].split(".")
        for p in parts[:-1]:
            cursor = cursor.setdefault(p, {})
        # A form newline is a new Word paragraph, not a soft line break.
        # docxtpl's paragraph separator clones the slot's paragraph/run
        # properties, including first-line indents, spacing and font.
        value = document_field_value(f)
        cursor[parts[-1]] = (
            value.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "\a")
        )
    context["participants"] = [{"name": r["name"]} for r in participants]
    # Keep the reference scoring sheet's three writing columns readable.
    # More units receive another complete pair of business/technical sheets.
    context["participant_groups"] = [
        dict(
            zip(("first", "second", "third"), names + [""] * (3 - len(names))),
        )
        for offset in range(0, len(participants), 3)
        for names in [[r["name"] for r in participants[offset : offset + 3]]]
    ]
    if recipient is not None:
        start, end = document_date(recipient["start"]), document_date(
            recipient["end"],
        )
        context.setdefault("invitation", {}).update(
            recipient=recipient["name"],
            start=start,
            end=end,
        )
        context.setdefault("events", {})["invitation_time"] = document_date(
            recipient["start"] + " 至 " + recipient["end"],
        )
    template = DocxTemplate(io.BytesIO(content))
    template.render(context, jinja_env=environment(), autoescape=True)
    output = io.BytesIO()
    template.save(output)
    # Catch placeholders accidentally emitted as values, too. Never label such
    # an output as confirmed or leave partially generated files on disk.
    with ZipFile(io.BytesIO(output.getvalue())) as archive:
        for name in archive.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                # lxml is a C extension; runtime behavior is covered by tests.
                # pylint: disable-next=c-extension-no-member
                xml = etree.fromstring(archive.read(name))
                if re.search(r"{{|{%|{#", "".join(xml.itertext())):
                    raise ValueError(
                        "生成内容仍有模板标记，请检查填写值和 Word 模板",
                    )
    return output.getvalue()


def export_excel(case):
    book = Workbook()
    sheet = book.active
    sheet.title = "变量配置"
    sheet.append(
        ["变量名称", "中文名称", "值", "分组", "类型", "必填", "文书显示"],
    )
    for f in case["fields"]:
        sheet.append(
            [
                f["key"],
                f["label"],
                f["value"],
                f["group"],
                f["kind"],
                "是" if f["required"] else "否",
                document_field_value(f),
            ],
        )
    recipients = book.create_sheet("受邀单位")
    recipients.append(
        [
            "单位名称",
            "领取开始时间",
            "领取结束时间",
            "开始时间跟随首家",
            "结束时间跟随首家",
            "领取时段（文书）",
        ],
    )
    for r in recipient_defaults(case["recipients"]):
        recipients.append(
            [
                r["name"],
                r["start"],
                r["end"],
                "是" if r["follow_start"] else "否",
                "是" if r["follow_end"] else "否",
                document_date(r["start"] + " 至 " + r["end"]),
            ],
        )
    rules = book.create_sheet("校核规则")
    rules.append(["校核项", "开始变量", "结束变量", "允许相同"])
    for r in case["rules"]:
        rules.append(
            [r["label"], r["start"], r["end"], "是" if r["equal"] else "否"],
        )
    check = book.create_sheet("日期校核")
    check.append(["校核项", "结果", "说明"])
    for r in checks(case):
        check.append(
            [
                r["label"],
                {"ok": "通过", "missing": "待补充", "error": "冲突"}[
                    r["status"]
                ],
                r["detail"],
            ],
        )
    for s in book:
        s.freeze_panes = "A2"
        s.auto_filter.ref = s.dimensions
        for cell in s[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="264F52")
        for col in s.columns:
            s.column_dimensions[col[0].column_letter].width = min(
                65,
                max(20, max(len(text(c.value)) for c in col) + 4),
            )
        # Preserve literal values beginning with =,+,-,@; never create formulas
        # from user-provided content.
        for row in list(s.rows)[1:]:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def starter_files():
    fields = [
        ("project.name", "项目名称", "", "项目资料", "text"),
        ("project.code", "项目编号", "", "项目资料", "text"),
        ("organization.name", "邀请单位", "", "单位与联系", "text"),
        ("contact.name", "联系人", "", "单位与联系", "text"),
        ("contact.phone", "联系电话", "", "单位与联系", "text"),
        ("dates.issue", "文件落款日期", "", "日期安排", "date"),
        (
            "events.qa_deadline",
            "答疑截止时间",
            "",
            "日期安排",
            "datetime-local",
        ),
        (
            "events.response_deadline",
            "响应截止时间",
            "",
            "日期安排",
            "datetime-local",
        ),
        (
            "events.negotiation_time",
            "谈判时间",
            "",
            "日期安排",
            "datetime-local",
        ),
        ("events.location", "谈判地点", "", "日期安排", "text"),
        (
            "qualification.start_year",
            "资格资料起始年份",
            "",
            "资格要求",
            "text",
        ),
    ]
    case = {
        "fields": [
            dict(
                zip(("key", "label", "value", "group", "kind"), f),
                required=True,
            )
            for f in fields
        ],
        "recipients": [{"name": "", "start": "", "end": ""} for _ in range(3)],
        "rules": [],
        "templates": [],
    }
    from docx import Document

    outputs = {"变量管理表.xlsx": export_excel(case)}
    for name, paragraphs in {
        "谈判文件_示例模板.docx": [
            "商务谈判文件",
            "项目名称：{{ project.name }}",
            "项目编号：{{ project.code }}",
            "邀请单位：{{ organization.name }}",
            "资格要求：自{{ qualification.start_year }}年至今，类似业绩不少于1个。",
            "答疑截止：{{ events.qa_deadline }}",
            "响应截止：{{ events.response_deadline }}",
            "谈判时间：{{ events.negotiation_time }}",
            "谈判地点：{{ events.location }}",
            "联系人：{{ contact.name }}　电话：{{ contact.phone }}",
            "{{ organization.name }}",
            "{{ dates.issue }}",
        ],
        "邀请函_示例模板.docx": [
            "谈判邀请函",
            "{{ invitation.recipient }}：",
            "邀请贵单位参加“{{ project.name }}”（编号：{{ project.code }}）商务谈判。",
            "邀请函领取时间：{{ events.invitation_time }}",
            "答疑截止：{{ events.qa_deadline }}",
            "响应截止：{{ events.response_deadline }}",
            "谈判时间：{{ events.negotiation_time }}",
            "地点：{{ events.location }}",
            "联系人：{{ contact.name }}　电话：{{ contact.phone }}",
            "{{ organization.name }}",
            "{{ dates.issue }}",
        ],
    }.items():
        doc = Document()
        doc.add_heading(paragraphs[0], 0)
        for p in paragraphs[1:]:
            doc.add_paragraph(p)
        stream = io.BytesIO()
        doc.save(stream)
        outputs[name] = stream.getvalue()
    return outputs
