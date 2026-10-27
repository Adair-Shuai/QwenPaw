# -*- coding: utf-8 -*-
"""Project stages and private material history around existing Word forms."""

import copy
import hashlib
import io
import json
import mimetypes
import re
import tempfile
import uuid
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import quote
from zipfile import BadZipFile, ZIP_DEFLATED, ZipFile

from fastapi import Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, Field
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader

STAGES = (
    (
        "plan",
        "上报计划",
        "明确项目需求、预算与计划批复",
        (
            "科研计划填报表",
            "外协必要性说明",
            "计划批复或系统回执",
        ),
    ),
    (
        "selection",
        "选商论证",
        "整理选商依据、单位资质与论证意见",
        (
            "选商论证材料",
            "拟选单位资质材料",
            "论证意见及审批记录",
            "非招标方案填报表及系统回执",
        ),
    ),
    (
        "negotiation",
        "商务谈判",
        "准备谈判文书，留存评分与结果材料",
        (
            "服务选商实施方案",
            "商务谈判文件",
            "谈判邀请函",
            "投标报价一览表",
            "综合评分表",
            "综合评分汇总表",
            "谈判纪要及结果审批",
        ),
    ),
    (
        "kickoff",
        "开题论证",
        "衔接合同与任务要求，整理开题材料",
        (
            "合同及任务书",
            "开题报告",
            "开题论证意见",
        ),
    ),
    (
        "contract",
        "合同审查",
        "核对合同条款、服务范围、金额与审批结果",
        (
            "技术服务合同及附件",
            "合同审查意见及修改记录",
            "合同审批或系统回执",
            "双方签署版合同",
        ),
    ),
    (
        "acceptance",
        "结题验收",
        "对照成果与考核要求，整理验收材料",
        (
            "结题报告",
            "成果及考核指标对照材料",
            "验收意见及证明",
        ),
    ),
    (
        "payment",
        "归档付款",
        "整理归档目录与付款所需凭据",
        (
            "归档材料目录",
            "合同及验收凭据",
            "发票及付款申请",
            "付款凭证",
        ),
    ),
)
STAGE_IDS = {stage[0] for stage in STAGES}
STAGE_STATUS = {"pending", "in_progress", "submitted", "completed"}
MATERIAL_STATUS = {"pending", "preparing", "ready", "submitted", "done", "na"}
MATERIAL_LABELS = {
    "pending": "待准备",
    "preparing": "整理中",
    "ready": "待提交",
    "submitted": "已提交",
    "done": "已完成",
    "na": "不适用",
}
SHARED_PREFIXES = (
    "project.",
    "org.",
    "organization.",
    "contact.",
    "research.",
    "deliverables.",
    "metrics.",
    "workload.",
    "milestones.",
    "budget.",
    "funding.",
)
EXTENSIONS = {
    ".doc",
    ".docx",
    ".xls",
    ".xlsx",
    ".ppt",
    ".pptx",
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".zip",
    ".txt",
    ".csv",
}


def material_index(project, stage_id=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "材料目录"
    sheet.append(
        [
            "阶段",
            "材料名称",
            "状态",
            "必备",
            "文件版本数",
            "最新文件",
            "计划日期",
            "办理系统",
            "提交编号/回执",
            "材料说明",
        ],
    )
    for stage in project["stages"]:
        if stage_id and stage["id"] != stage_id:
            continue
        for material in stage["materials"]:
            sheet.append(
                [
                    stage["title"],
                    material["name"],
                    MATERIAL_LABELS[material["status"]],
                    "是" if material["required"] else "否",
                    len(material["files"]),
                    material["files"][-1]["name"] if material["files"] else "",
                    stage["due_date"],
                    stage["system_name"],
                    stage["submission_ref"],
                    material["note"],
                ],
            )
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2C6260")
    for column, width in zip(
        "ABCDEFGHIJ",
        (16, 34, 12, 8, 12, 42, 16, 25, 30, 45),
    ):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def now():
    return datetime.now(timezone.utc).isoformat()


def identifier(value):
    if not re.fullmatch(r"[a-f0-9]{32}", value):
        raise HTTPException(404, "项目或材料不存在")
    return value


def shared_fields(fields):
    return [
        copy.deepcopy(f)
        for f in fields
        if f["key"].startswith(SHARED_PREFIXES)
    ]


def values(fields):
    return {field["key"]: field["value"] for field in fields}


def content_digest(case):
    return hashlib.sha256(
        json.dumps(
            {
                k: case.get(k, [])
                for k in ("fields", "recipients", "rules", "templates")
            },
            sort_keys=True,
            ensure_ascii=False,
        ).encode(),
    ).hexdigest()


def default_stages():
    return [
        {
            "id": key,
            "title": title,
            "purpose": purpose,
            "status": "pending",
            "notes": "",
            "due_date": "",
            "system_name": "",
            "system_url": "",
            "submission_ref": "",
            "content_revision": 0,
            "acknowledged": None,
            "completed_at": None,
            "materials": [
                {
                    "id": uuid.uuid4().hex,
                    "name": name,
                    "required": False,
                    "status": "pending",
                    "note": "",
                    "files": [],
                }
                for name in materials
            ],
        }
        for key, title, purpose, materials in STAGES
    ]


class Revision(BaseModel):
    revision: int = Field(ge=1)


class ProjectBody(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    code: str = Field(default="", max_length=100)
    owner: str = Field(default="", max_length=100)


class ProfileBody(ProjectBody, Revision):
    values: dict[str, str] = Field(default_factory=dict, max_length=200)


class StageBody(Revision):
    status: str
    confirmed: bool = False
    notes: str = Field(default="", max_length=5000)
    due_date: str = Field(default="", max_length=10)
    system_name: str = Field(default="", max_length=150)
    system_url: str = Field(default="", max_length=1000)
    submission_ref: str = Field(default="", max_length=500)


class MaterialBody(Revision):
    name: str = Field(min_length=1, max_length=150)
    required: bool = False
    status: str = "pending"
    note: str = Field(default="", max_length=3000)


class CaseBody(Revision):
    stage_id: str


class SyncBody(Revision):
    case_revision: int
    keys: list[str] = Field(min_length=1, max_length=200)


class ApprovalApplyBody(Revision):
    file_id: str = Field(min_length=32, max_length=32)
    keys: list[str] = Field(min_length=1, max_length=20)


APPROVAL_FIELDS = (
    ("project.executing_unit", "承担单位"),
    ("project.name", "项目名称"),
    ("plan.budget_wan", "外协项目经费（万元）"),
    ("plan.category", "类别"),
    ("research.contents", "研究内容"),
    ("funding.project_name", "依托项目（课题）名称"),
    ("funding.amount_wan", "依托项目经费（万元）"),
    ("funding.wbs", "WBS 编码"),
)


def _join_pdf_parts(parts):
    return "".join(text.strip() for _, _, text in parts if text.strip())


def approval_from_pdf(content):
    """Read the institute's plan approval table without model inference."""
    try:
        page = PdfReader(io.BytesIO(content)).pages[0]
        parts = []

        def collect(text, _cm, tm, _font, _size):
            if text.strip():
                parts.append((float(tm[5]), float(tm[4]), text.strip()))

        # PyPDF exposes these methods dynamically on its page proxy.
        # pylint: disable=no-member
        plain = page.extract_text(visitor_text=collect) or ""
        layout = page.extract_text(extraction_mode="layout") or plain
        # pylint: enable=no-member
    except Exception as exc:  # pragma: no cover - parser details vary by PDF
        return {
            "parser_version": 1,
            "status": "unreadable",
            "message": f"审批文件无法解析：{exc}",
        }
    date_match = re.search(r"批复时间[：:]\s*(\d{4})年(\d{1,2})月(\d{1,2})日", plain)
    unit_match = re.search(r"单位名称[：:]\s*([^\n]+?)(?:\s+批复时间|$)", plain)
    if not date_match or not unit_match or "研究内容" not in plain:
        return {
            "parser_version": 1,
            "status": "unsupported",
            "message": "未识别到标准计划审批表，请人工核对审批结果。",
        }
    body = [item for item in parts if 330 <= item[0] <= 410]
    columns = {
        "project.name": (60, 160),
        "plan.budget_wan": (160, 250),
        "plan.category": (250, 320),
        "research.contents": (320, 500),
        "funding.project_name": (500, 646),
        "funding.amount_wan": (646, 720),
        "funding.wbs": (720, 820),
    }
    extracted = {
        key: _join_pdf_parts(
            sorted(
                (item for item in body if low <= item[1] < high),
                key=lambda item: (-item[0], item[1]),
            ),
        )
        for key, (low, high) in columns.items()
    }
    extracted["project.executing_unit"] = unit_match.group(1).strip()
    # Keep numbered research items readable in the form and generated files.
    extracted["research.contents"] = re.sub(
        r"(?<!^)\s*(?=\d+[.、])",
        "\n",
        extracted["research.contents"],
    )
    extracted["funding.wbs"] = re.sub(r"\s+", "", extracted["funding.wbs"])
    labels = dict(APPROVAL_FIELDS)
    fields = [
        {"key": key, "label": labels[key], "value": extracted[key]}
        for key, _label in APPROVAL_FIELDS
        if extracted.get(key)
    ]
    if len(fields) < 5:
        return {
            "parser_version": 1,
            "status": "unsupported",
            "message": "审批表字段不完整，请人工核对审批结果。",
        }
    return {
        "parser_version": 1,
        "status": "ready",
        "approved_at": "-".join(
            f"{int(value):02d}" if index else value
            for index, value in enumerate(date_match.groups())
        ),
        "title": next(
            (line.strip() for line in layout.splitlines() if "实施计划" in line),
            "计划审批结果",
        ),
        "fields": fields,
    }


class _ProjectStorage:
    """Atomic project metadata and shared stage/material addressing."""

    def __init__(self, host):
        self.host = host

    @staticmethod
    def path(root, project_id):
        return root / "projects" / identifier(project_id)

    def load(self, root, project_id):
        path = self.path(root, project_id) / "project.json"
        if not path.is_file():
            raise HTTPException(404, "项目不存在")
        project = json.loads(path.read_text(encoding="utf-8"))
        migrated = False
        current_stages = {stage["id"]: stage for stage in project["stages"]}
        if set(current_stages) != STAGE_IDS:
            project["stages"] = [
                current_stages.get(stage["id"], stage)
                for stage in default_stages()
            ]
            migrated = True
        negotiation = next(
            s for s in project["stages"] if s["id"] == "negotiation"
        )
        if not any(m["name"] == "服务选商实施方案" for m in negotiation["materials"]):
            default = next(
                s for s in default_stages() if s["id"] == "negotiation"
            )
            negotiation["materials"].insert(0, default["materials"][0])
            migrated = True
        for stage in project["stages"]:
            if stage["id"] != "plan":
                continue
            for material in stage["materials"]:
                for file in material["files"]:
                    if (
                        not file.get("approval")
                        or file["approval"].get("parser_version") == 1
                    ):
                        continue
                    source = (
                        self.path(root, project_id) / "files" / file["storage"]
                    )
                    if source.is_file():
                        file["approval"] = approval_from_pdf(
                            source.read_bytes(),
                        )
                        migrated = True
        existing = {f["key"] for f in project["fields"]}
        project["fields"].extend(
            f
            for f in self.host.builtin.additional_shared_fields()
            if f["key"] not in existing
        )
        if migrated:
            self.save(root, project, update_time=False)
        return project

    def save(self, root, project, update_time=True):
        path = self.path(root, project["id"])
        path.mkdir(parents=True, exist_ok=True)
        if update_time:
            project["updated_at"] = now()
        temp = path / "project.tmp"
        temp.write_text(
            json.dumps(project, ensure_ascii=False),
            encoding="utf-8",
        )
        temp.replace(path / "project.json")

    @staticmethod
    def check(project, revision):
        if project["revision"] != revision:
            raise HTTPException(409, "项目已在其他页面修改，请重新打开后编辑")

    @staticmethod
    def stage(project, stage_id):
        if stage_id not in STAGE_IDS:
            raise HTTPException(404, "项目阶段不存在")
        return next(s for s in project["stages"] if s["id"] == stage_id)

    def material(self, project, stage_id, material_id):
        stage = self.stage(project, stage_id)
        result = next(
            (m for m in stage["materials"] if m["id"] == material_id),
            None,
        )
        if not result:
            raise HTTPException(404, "材料不存在")
        return stage, result

    def touch(self, project, stage, root):
        if stage["acknowledged"] is None:
            stage["acknowledged"] = self.dependencies(
                project,
                stage["id"],
                root,
            )
        stage["content_revision"] += 1
        if stage["status"] in {"pending", "completed"}:
            stage["status"] = "in_progress"
        stage["completed_at"] = None
        project["revision"] += 1


class _ProjectForms(_ProjectStorage):
    """Form preparation, linked stage content and field synchronization."""

    def form_fingerprint(self, root, project, stage_id):
        states = []
        for link in project["forms"]:
            if link["case_id"] == project.get("form_id"):
                case = self.host.load(root, link["case_id"])
                context = case.get("stage_contexts", {}).get(stage_id)
                if case.get("stage_id") == stage_id or context:
                    states.append(
                        (
                            case["id"],
                            content_digest(
                                (
                                    case
                                    if case.get("stage_id") == stage_id
                                    else context
                                ),
                            ),
                            [
                                g["id"]
                                for g in case["generations"]
                                if g.get(
                                    "stage_id",
                                    case.get("original_stage_id"),
                                )
                                == stage_id
                            ],
                        ),
                    )
                continue
            if link["stage_id"] == stage_id:
                case = self.host.load(root, link["case_id"])
                states.append(
                    (
                        case["id"],
                        case["revision"],
                        [g["id"] for g in case["generations"]],
                    ),
                )
        return hashlib.sha256(
            json.dumps(states, sort_keys=True).encode(),
        ).hexdigest()

    def dependencies(self, project, stage_id, root):
        prior = {}
        for stage in project["stages"]:
            if stage["id"] == stage_id:
                break
            prior[stage["id"]] = [
                stage["content_revision"],
                self.form_fingerprint(root, project, stage["id"]),
                project["shared_revision"],
            ]
        return prior

    def form_summary(self, root, project, link):
        case = self.host.load(root, link["case_id"])
        current = values(project["fields"])
        snapshot = case.get(
            "project_snapshot",
            values(shared_fields(case["fields"])),
        )
        changed = [
            f["key"]
            for f in case["fields"]
            if f["key"] in current
            and snapshot.get(f["key"]) != current[f["key"]]
        ]
        return {
            "case_id": case["id"],
            "original_stage_id": case.get(
                "original_stage_id",
                link["stage_id"],
            ),
            "stage_id": (
                case.get("stage_id", link["stage_id"])
                if case.get("project_form")
                else link["stage_id"]
            ),
            "name": case["name"],
            "revision": case["revision"],
            "updated_at": case["updated_at"],
            "outputs": sum(len(g["files"]) for g in case["generations"]),
            "shared_changes": changed,
            "generations": [
                {k: v for k, v in g.items() if k != "review_token"}
                for g in case["generations"]
            ],
        }

    def open_form(self, root, project, stage_id, workbook=None):
        """One workbook per project; stages are pages, never new runs."""
        # pylint: disable=too-many-branches,too-many-statements
        self.stage(project, stage_id)
        imported = None
        if workbook:
            try:
                imported = self.host.read_excel(
                    self.host.upload(workbook, ".xlsx"),
                )
            except Exception as exc:
                if isinstance(exc, HTTPException):
                    raise
                raise HTTPException(422, f"Excel 无法读取：{exc}") from exc
        keys = (
            "fields",
            "recipients",
            "rules",
            "templates",
            "project_snapshot",
            "template_notice",
        )
        case_id = project.get("form_id")
        if not case_id:
            candidates = [
                self.host.load(root, f["case_id"]) for f in project["forms"]
            ]
            candidates.sort(
                key=lambda c: (
                    c.get("stage_id") == "negotiation",
                    c["updated_at"],
                ),
                reverse=True,
            )
            created = (
                candidates[0]
                if candidates
                else self.add_case(root, project, stage_id)
            )
            case = self.host.load(root, created["id"])
            # Preserve the original bytes and old confirmations.
            backup = self.path(root, project["id"]) / "form-history"
            backup.mkdir(exist_ok=True)
            (backup / (case["id"] + "_original.json")).write_bytes(
                (
                    self.host.folder(root, case["id"]) / "draft.json"
                ).read_bytes(),
            )
            project["form_id"] = case["id"]
            project["revision"] += 1
            self.save(root, project)
        else:
            case = self.host.load(root, case_id)
        case.setdefault(
            "original_stage_id",
            case.get("stage_id", "negotiation"),
        )
        contexts = case.setdefault("stage_contexts", {})
        new_page = stage_id not in contexts
        current = case.get("stage_id", "negotiation")
        if stage_id == "contract" and workbook:
            raise HTTPException(422, "合同审查直接引用前序资料，请在原阶段修改")
        if current == "contract":
            self.prepare_contract(root, project, case)
        if current == "plan":
            self.prepare_plan(root, project, case)
        if current == "selection":
            self.prepare_selection(root, project, case)
        if current == "acceptance":
            self.prepare_acceptance(root, project, case)
        if current == "negotiation":
            self.prepare_implementation(root, project, case)
        if (
            current == stage_id
            and workbook is None
            and case.get("project_form")
        ):
            return self.host.public(case)
        common = values(shared_fields(case["fields"]))
        contexts[current] = {
            k: copy.deepcopy(case[k]) for k in keys if k in case
        }
        contexts[current]["revision"] = case["revision"]
        if workbook:
            history = self.path(root, project["id"]) / "form-history"
            history.mkdir(exist_ok=True)
            (history / f"{case['id']}_r{case['revision']}.json").write_text(
                json.dumps(case, ensure_ascii=False),
                encoding="utf-8",
            )
            data = imported
            data["templates"] = contexts.get(stage_id, {}).get("templates", [])
        elif stage_id in contexts:
            data = copy.deepcopy(contexts[stage_id])
        else:
            data = {
                "fields": (
                    self.host.builtin.blank_fields()
                    if stage_id == "negotiation"
                    else (
                        self.host.builtin.plan_fields(project["fields"])
                        if stage_id == "plan"
                        else (
                            self.host.builtin.selection_fields(
                                project["fields"],
                            )
                            if stage_id == "selection"
                            else (
                                self.host.builtin.acceptance_fields(
                                    project["fields"],
                                )
                                if stage_id == "acceptance"
                                else copy.deepcopy(project["fields"])
                            )
                        )
                    )
                ),
                "recipients": (
                    [{"name": "", "start": "", "end": ""} for _ in range(3)]
                    if stage_id == "negotiation"
                    else []
                ),
                "rules": [],
                "templates": [],
            }
        for key in keys:
            case.pop(key, None)
        case.update({k: copy.deepcopy(data[k]) for k in keys if k in data})
        for field in case["fields"]:
            if not workbook and field["key"] in common:
                field["value"] = common[field["key"]]
            elif not field["value"]:
                field["value"] = common.get(
                    field["key"],
                    values(project["fields"]).get(field["key"], ""),
                )
        case.update(
            stage_id=stage_id,
            name="项目表单",
            project_form=True,
            review_token=None,
            revision=case["revision"] + 1,
        )
        case.setdefault(
            "project_snapshot",
            {
                f["key"]: values(project["fields"]).get(f["key"], "")
                for f in case["fields"]
            },
        )
        if stage_id == "plan":
            self.prepare_plan(root, project, case)
        if stage_id == "selection":
            self.prepare_selection(root, project, case)
        if stage_id == "acceptance":
            self.prepare_acceptance(root, project, case)
        if stage_id == "negotiation":
            self.prepare_implementation(root, project, case)
        if stage_id == "contract":
            self.prepare_contract(root, project, case)
        if not case["templates"]:
            try:
                self.host.attach_default_templates(root, case)
            except (ValueError, KeyError, OSError):
                case["template_notice"] = "默认模板未能载入，请重新选择 Word 文件"
        self.host.save(root, case)
        if workbook or (
            new_page and self.stage(project, stage_id)["status"] == "pending"
        ):
            self.touch(project, self.stage(project, stage_id), root)
        self.save(root, project)
        return self.host.public(case)

    def prepare_contract(self, root, project, case):
        """Contract review references upstream values rather than \
another copy."""
        references = {f["key"]: copy.deepcopy(f) for f in project["fields"]}
        contexts = case.get("stage_contexts", {})
        for stage_id in ("negotiation", "selection", "plan"):
            for field in contexts.get(stage_id, {}).get("fields", []):
                if field["key"] not in references:
                    references[field["key"]] = copy.deepcopy(field)
        for field in references.values():
            field.update(scope="shared", read_only=True, required=False)
        fields = list(references.values())
        if fields == case["fields"] and case.get("contract_schema") == 1:
            return
        history = self.path(root, project["id"]) / "form-history"
        history.mkdir(exist_ok=True)
        source = self.host.folder(root, case["id"]) / "draft.json"
        if source.is_file():
            backup = history / (
                f"{case['id']}"
                f"_before_contract_references_r"
                f"{case['revision']}"
                f".json"
            )
            if not backup.exists():
                backup.write_bytes(source.read_bytes())
        case["fields"] = fields
        case["recipients"] = []
        case["contract_schema"] = 1
        case["review_token"] = None
        case["revision"] += 1
        self.host.save(root, case)

    def prepare_implementation(self, root, project, case):
        """Extend bundled negotiation pages while retaining custom \
templates."""
        if case.get("implementation_schema") == 1:
            return
        existing = {field["key"] for field in case["fields"]}
        if (
            not {
                "project.name",
                "contact.department",
                "place.negotiation",
                "events.negotiation_time",
                "dates.issue",
                "qualification.start_year",
            }
            <= existing
        ):
            return
        history = self.path(root, project["id"]) / "form-history"
        history.mkdir(exist_ok=True)
        source = self.host.folder(root, case["id"]) / "draft.json"
        backup = (
            history
            / f"{case['id']}_before_implementation_r{case['revision']}.json"
        )
        if source.is_file() and not backup.exists():
            backup.write_bytes(source.read_bytes())
        case["fields"].extend(
            field
            for field in self.host.builtin.implementation_fields()
            if field["key"] not in existing
        )
        name = self.host.builtin.NAMES[-1]
        if not any(template["name"] == name for template in case["templates"]):
            if not case["templates"]:
                self.host.attach_default_templates(root, case)
            else:
                content = self.host.builtin.load_templates()[name]
                info = self.host.template_info(
                    content,
                    name,
                    {field["key"] for field in case["fields"]},
                )
                info.update(
                    id=uuid.uuid4().hex,
                    sha256=hashlib.sha256(content).hexdigest(),
                    from_default=True,
                    from_builtin=True,
                )
                target = self.host.folder(root, case["id"]) / "templates"
                target.mkdir(exist_ok=True)
                (target / (info["id"] + ".docx")).write_bytes(content)
                case["templates"].append(info)
        case["implementation_schema"] = 1
        case["review_token"] = None
        case["revision"] += 1
        self.host.save(root, case)

    def prepare_selection(self, root, project, case):
        if case.get("selection_schema") == 1:
            self.sync_selection_fields(root, project, case)
            return
        fixed = self.host.builtin.selection_fields(project["fields"])
        old = values(case["fields"])
        # Custom forms keep their own schema and templates.
        if case["templates"] and not any(
            t.get("from_builtin") for t in case["templates"]
        ):
            return
        history = self.path(root, project["id"]) / "form-history"
        history.mkdir(exist_ok=True)
        source = self.host.folder(root, case["id"]) / "draft.json"
        backup = (
            history / f"{case['id']}_before_selection_r{case['revision']}.json"
        )
        if source.is_file() and not backup.exists():
            backup.write_bytes(source.read_bytes())
        for field in fixed:
            if field["key"] in old and not field.get("read_only"):
                field["value"] = old[field["key"]]
        case["fields"] = fixed
        case["selection_schema"] = 1
        case["review_token"] = None
        case["revision"] += 1
        if not case["templates"]:
            self.host.attach_default_templates(root, case)
        self.host.save(root, case)

    def prepare_acceptance(self, root, project, case):
        """Upgrade pre-template acceptance pages to the focused opinion \
form."""
        if case.get("acceptance_schema") == 1:
            return
        if case["templates"] and not any(
            template.get("from_builtin") for template in case["templates"]
        ):
            return
        fixed = self.host.builtin.acceptance_fields(project["fields"])
        old = values(case["fields"])
        history = self.path(root, project["id"]) / "form-history"
        history.mkdir(exist_ok=True)
        source = self.host.folder(root, case["id"]) / "draft.json"
        backup = (
            history
            / f"{case['id']}_before_acceptance_r{case['revision']}.json"
        )
        if source.is_file() and not backup.exists():
            backup.write_bytes(source.read_bytes())
        for field in fixed:
            if field["key"] in old:
                field["value"] = old[field["key"]]
        case["fields"] = fixed
        case["recipients"] = []
        case["acceptance_schema"] = 1
        case["review_token"] = None
        case["revision"] += 1
        if not case["templates"]:
            self.host.attach_default_templates(root, case)
        self.host.save(root, case)

    def sync_selection_fields(self, root, project, case):
        """Refresh protected fields without touching historical generations."""
        definitions = {
            f["key"]: f
            for f in self.host.builtin.selection_fields(project["fields"])
        }
        before = copy.deepcopy(case["fields"])
        for field in case["fields"]:
            definition = definitions.get(field["key"], {})
            if "options" in definition:
                field["options"] = definition["options"]
            if definition.get("read_only"):
                field.update(
                    {
                        key: definition[key]
                        for key in ("read_only", "value", "help")
                    },
                )
        if before == case["fields"]:
            return
        history = self.path(root, project["id"]) / "form-history"
        history.mkdir(exist_ok=True)
        source = self.host.folder(root, case["id"]) / "draft.json"
        backup = history / (
            f"{case['id']}"
            f"_before_selection_controls_r"
            f"{case['revision']}"
            f".json"
        )
        if source.is_file() and not backup.exists():
            backup.write_bytes(source.read_bytes())
        case["revision"] += 1
        case["review_token"] = None
        self.host.save(root, case)

    def prepare_plan(self, root, project, case):
        if (
            case["templates"]
            and not any(t.get("from_builtin") for t in case["templates"])
            and not {"plan.name", "plan.budget_wan", "plan.necessity"}
            <= set(values(case["fields"]))
        ):
            case.pop("plan_schema", None)
            return
        references = {
            f["key"]: f
            for f in self.host.builtin.plan_fields(project["fields"])
            if f.get("scope") == "shared"
        }
        for field in case["fields"]:
            if field["key"] in references:
                field["scope"] = "shared"
        if case.get("plan_schema") == 1:
            self.prepare_full_necessity(root, project, case)
            return
        history = self.path(root, project["id"]) / "form-history"
        history.mkdir(exist_ok=True)
        source = self.host.folder(root, case["id"]) / "draft.json"
        if source.is_file():
            backup = (
                history / f"{case['id']}_before_plan_r{case['revision']}.json"
            )
            if not backup.exists():
                backup.write_bytes(source.read_bytes())
        old = values(case["fields"])
        fixed = self.host.builtin.plan_fields(project["fields"])
        if not case["templates"]:
            for field in fixed:
                if old.get(field["key"]):
                    field["value"] = old[field["key"]]
            case["fields"] = fixed
        else:
            existing = {f["key"] for f in case["fields"]}
            case["fields"].extend(f for f in fixed if f["key"] not in existing)
        case["plan_schema"] = 1
        case["review_token"] = None
        case["revision"] += 1
        if not case["templates"]:
            self.host.attach_default_templates(root, case)
        self.host.save(root, case)

    def prepare_full_necessity(self, root, project, case):
        """Upgrade the exact old template while retaining its output."""
        legacy = self.host.builtin.LEGACY_PLAN_WORD
        old_templates = [t for t in case["templates"] if t["sha256"] == legacy]
        if not old_templates:
            return
        target = self.host.folder(root, case["id"])
        for template in old_templates:
            content = (
                target / "templates" / (template["id"] + ".docx")
            ).read_bytes()
            if hashlib.sha256(content).hexdigest() != legacy:
                raise HTTPException(422, "原必要性说明模板已改变，请重新核对模板")
        history = self.path(root, project["id"]) / "form-history"
        history.mkdir(exist_ok=True)
        backup = (
            history
            / f"{case['id']}_before_full_necessity_r{case['revision']}.json"
        )
        if not backup.exists():
            backup.write_bytes((target / "draft.json").read_bytes())
        content = self.host.builtin.load_templates("plan")[
            self.host.builtin.PLAN_NAMES[0]
        ]
        for template in old_templates:
            template_id = uuid.uuid4().hex
            info = self.host.template_info(content, template["name"])
            (target / "templates" / (template_id + ".docx")).write_bytes(
                content,
            )
            template.update(
                info,
                id=template_id,
                sha256=hashlib.sha256(content).hexdigest(),
            )
        fields = {f["key"]: f for f in case["fields"]}
        necessity = fields.get("plan.necessity")
        background = fields.get("project.background", {}).get("value", "")
        if necessity:
            text = necessity["value"]
            if text.strip() and background and not text.startswith(background):
                necessity["value"] = background + text
            definition = next(
                f
                for f in self.host.builtin.plan_fields()
                if f["key"] == "plan.necessity"
            )
            necessity["help"] = definition["help"]
        if not any(
            "project.background" in t["keys"] for t in case["templates"]
        ):
            case["fields"] = [
                f for f in case["fields"] if f["key"] != "project.background"
            ]
        case["revision"] += 1
        case["review_token"] = None
        self.host.save(root, case)

    def sync(self, root, project, case_id, body):
        link = self.linked(project, case_id)
        case = self.host.load(root, case_id)
        self.host.version(case, body.case_revision)
        current = values(project["fields"])
        allowed = set(current) & set(values(case["fields"]))
        if set(body.keys) - allowed:
            raise HTTPException(422, "所选字段不属于项目共用资料")
        snapshot = case.get(
            "project_snapshot",
            values(shared_fields(case["fields"])),
        )
        for field in case["fields"]:
            if field["key"] in body.keys:
                field["value"] = current[field["key"]]
                snapshot[field["key"]] = current[field["key"]]
        case["project_snapshot"] = snapshot
        case["revision"] += 1
        case["review_token"] = None
        self.host.save(root, case)
        self.touch(
            project,
            self.stage(
                project,
                (
                    case.get("stage_id", link["stage_id"])
                    if case.get("project_form")
                    else link["stage_id"]
                ),
            ),
            root,
        )
        self.save(root, project)
        return self.host.public(case)


class _ProjectArchive(_ProjectStorage):
    """Immutable project archive inventory, attachments and rendered forms."""

    def _archive_manifest(self, root, project, stage_id):
        manifest = self.public(root, project)
        if stage_id:
            manifest["stages"] = [
                s for s in manifest["stages"] if s["id"] == stage_id
            ]
            manifest["forms"] = [
                f for f in manifest["forms"] if f["stage_id"] == stage_id
            ]
        for stage in manifest["stages"]:
            for material in stage["materials"]:
                for file in material["files"]:
                    file.pop("preview_url", None)
        for form in manifest["forms"]:
            for generation in form["generations"]:
                for file in generation["files"]:
                    file.pop("preview_url", None)
        return json.dumps(manifest, ensure_ascii=False, indent=2)

    def _archive_attachments(self, archive, root, project, titles, stage_id):
        for stage in project["stages"]:
            if stage_id and stage["id"] != stage_id:
                continue
            for number, material in enumerate(stage["materials"], 1):
                title = self.host.name_only(material["name"])[:60]
                entry = f"{titles[stage['id']]}/{number:02d}_{title}"
                for index, file in enumerate(material["files"], 1):
                    source = (
                        self.path(root, project["id"])
                        / "files"
                        / file["storage"]
                    )
                    if (
                        not source.is_file()
                        or hashlib.sha256(source.read_bytes()).hexdigest()
                        != file["sha256"]
                    ):
                        raise HTTPException(422, "材料文件缺失或校验失败，请补充后再打包")
                    archive.write(
                        source,
                        entry + f"/V{index:02d}_{file['name']}",
                    )

    def _archive_forms(self, archive, root, project, titles, stage_id):
        # pylint: disable=too-many-branches
        for number, link in enumerate(project["forms"], 1):
            canonical = link["case_id"] == project.get("form_id")
            if stage_id and link["stage_id"] != stage_id and not canonical:
                continue
            case = self.host.load(root, link["case_id"])
            if canonical:
                contexts = copy.deepcopy(case.get("stage_contexts", {}))
                contexts[case["stage_id"]] = case
                for page_id, context in contexts.items():
                    if stage_id and page_id != stage_id:
                        continue
                    page = dict(
                        case,
                        **{
                            k: context[k]
                            for k in (
                                "fields",
                                "recipients",
                                "rules",
                                "templates",
                            )
                        },
                    )
                    prefix = titles[page_id] + "/项目表单/"
                    archive.writestr(
                        prefix + "填写数据.xlsx",
                        self.host.export_excel(page),
                    )
                    for t in page["templates"]:
                        source = (
                            self.host.folder(root, case["id"])
                            / "templates"
                            / (t["id"] + self.host.template_extension(t))
                        )
                        if (
                            not source.is_file()
                            or hashlib.sha256(source.read_bytes()).hexdigest()
                            != t["sha256"]
                        ):
                            raise HTTPException(422, "项目表单模板缺失或校验失败")
                        archive.write(
                            source,
                            prefix
                            + ("文书模板/" if page_id == "plan" else "Word模板/")
                            + self.host.name_only(t["name"]),
                        )
                if not stage_id:
                    for history in (
                        self.path(root, project["id"]) / "form-history"
                    ).glob("*.json"):
                        archive.write(history, "项目表单/历史底稿/" + history.name)
                    for source in (
                        self.host.folder(root, case["id"]) / "templates"
                    ).glob("*"):
                        archive.write(source, "项目表单/模板原件/" + source.name)
            title = f"{number:02d}_" + self.host.name_only(case["name"])[:45]
            form_path = titles[link["stage_id"]] + "/表单底稿/" + title
            if not canonical:
                archive.writestr(
                    form_path + "/填写数据.xlsx",
                    self.host.export_excel(case),
                )
            for index, template in enumerate(
                [] if canonical else case["templates"],
                1,
            ):
                source = (
                    self.host.folder(root, case["id"])
                    / "templates"
                    / (template["id"] + self.host.template_extension(template))
                )
                if (
                    not source.is_file()
                    or hashlib.sha256(source.read_bytes()).hexdigest()
                    != template["sha256"]
                ):
                    raise HTTPException(422, "表单模板缺失或校验失败，请检查后再打包")
                target = form_path + f"/Word模板/{index:02d}_"
                archive.write(
                    source,
                    target + self.host.name_only(template["name"]),
                )
            for index, generation in enumerate(case["generations"], 1):
                output_stage = generation.get(
                    "stage_id",
                    case.get("original_stage_id", link["stage_id"]),
                )
                if canonical and stage_id and output_stage != stage_id:
                    continue
                source = (
                    self.host.folder(root, case["id"])
                    / "outputs"
                    / generation["id"]
                    / "全部文书.zip"
                )
                try:
                    with ZipFile(source) as bundle:
                        names = bundle.namelist()
                        valid = (
                            bundle.testzip() is None
                            and len(names) == len(set(names))
                            and {f["name"] for f in generation["files"]}
                            <= set(names)
                        )
                        if not valid:
                            raise ValueError("Invalid generated archive")
                except (OSError, BadZipFile, ValueError) as exc:
                    raise HTTPException(422, "生成批次缺失或损坏，请检查后再打包") from exc
                target = titles[output_stage] + "/生成文书/" + title
                archive.write(source, target + f"_批次{index:02d}.zip")

    def archive(self, root, project, stage_id=None):
        # The response generator closes this stream after transmission.
        # pylint: disable-next=consider-using-with
        stream = tempfile.SpooledTemporaryFile(max_size=16 * 1024 * 1024)
        titles = {
            s["id"]: f"{i:02d}_{s['title']}"
            for i, s in enumerate(project["stages"], 1)
        }
        try:
            with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
                archive.writestr(
                    "项目材料台账.json",
                    self._archive_manifest(root, project, stage_id),
                )
                archive.writestr(
                    "材料目录.xlsx",
                    material_index(project, stage_id),
                )
                self._archive_attachments(
                    archive,
                    root,
                    project,
                    titles,
                    stage_id,
                )
                self._archive_forms(archive, root, project, titles, stage_id)
                from .review import (  # pylint: disable=E0402
                    archive_reviews,
                    archive_assistant,
                )

                archive_assistant(
                    archive,
                    self.path(root, project["id"]),
                    stage_id,
                )
                if stage_id in (None, "negotiation"):
                    archive_reviews(archive, self.path(root, project["id"]))
        except Exception:
            stream.close()
            raise
        stream.seek(0)
        return stream


class ProjectStore(_ProjectForms, _ProjectArchive):
    """JSON metadata is committed atomically; attachments are immutable."""

    def create(self, root, body, fields=None, persist=True):
        fields = shared_fields(fields or self.host.builtin.blank_fields())
        existing = {f["key"] for f in fields}
        fields.extend(
            f
            for f in self.host.builtin.additional_shared_fields()
            if f["key"] not in existing
        )
        if not any(f["key"] == "project.code" for f in fields):
            fields.append(
                {
                    "key": "project.code",
                    "label": "项目编号",
                    "group": "项目资料",
                    "kind": "text",
                    "required": False,
                    "value": body.code.strip(),
                },
            )
        for field in fields:
            if field["key"] == "project.name":
                field["value"] = body.name.strip()
            if field["key"] == "project.code":
                field["value"] = body.code.strip()
        project = {
            "id": uuid.uuid4().hex,
            "name": body.name.strip(),
            "code": body.code.strip(),
            "owner": body.owner.strip(),
            "fields": fields,
            "shared_revision": 1,
            "revision": 1,
            "created_at": now(),
            "stages": default_stages(),
            "forms": [],
        }
        if persist:
            self.save(root, project)
        return project

    def all_projects(self, root):
        return [
            json.loads(p.read_text(encoding="utf-8"))
            for p in root.glob("projects/*/project.json")
        ]

    def migrate(self, root):
        """Index legacy drafts without rewriting their bytes or outputs."""
        projects = self.all_projects(root)
        linked = {f["case_id"] for p in projects for f in p["forms"]}
        drafts = [
            json.loads(p.read_text(encoding="utf-8"))
            for p in root.glob("*/draft.json")
        ]
        for case in sorted(
            drafts,
            key=lambda c: c["updated_at"],
            reverse=True,
        ):
            if case["id"] in linked:
                continue
            data = values(case["fields"])
            name = data.get("project.name", "").strip() or case["name"]
            code = data.get("project.code", "").strip()
            identity = "code:" + code if code else "name:" + name
            project = next(
                (
                    p
                    for p in projects
                    if p["id"] == case.get("project_id")
                    or p.get("legacy_identity") == identity
                ),
                None,
            )
            if project is None:
                project = self.create(
                    root,
                    ProjectBody(name=name, code=code),
                    case["fields"],
                    persist=False,
                )
                project["legacy_identity"] = identity
                project["migrated"] = True
                project["created_at"] = case["updated_at"]
                project["updated_at"] = case["updated_at"]
                projects.append(project)
            stage_id = case.get("stage_id", "negotiation")
            project["forms"].append(
                {"case_id": case["id"], "stage_id": stage_id},
            )
            stage = self.stage(project, stage_id)
            stage["status"] = "in_progress"
            stage["content_revision"] += 1
            stage["acknowledged"] = self.dependencies(project, stage_id, root)
            project["updated_at"] = max(
                project["updated_at"],
                case["updated_at"],
            )
            self.save(root, project, update_time=False)
            linked.add(case["id"])
        for project in projects:
            if project.get("migrated") and not project.get(
                "legacy_order_version",
            ):
                if project["revision"] == 1:
                    ids = {f["case_id"] for f in project["forms"]}
                    dates = [c["updated_at"] for c in drafts if c["id"] in ids]
                    if dates:
                        project["updated_at"] = max(dates)
                project["legacy_order_version"] = 1
                self.save(root, project, update_time=False)
        return projects

    def public(self, root, project):
        result = copy.deepcopy(project)
        result["single_form"] = True
        result["forms"] = [
            self.form_summary(root, project, f) for f in project["forms"]
        ]
        result["forms"].sort(key=lambda f: f["updated_at"], reverse=True)
        for stage in result["stages"]:
            expected = stage["acknowledged"]
            current = self.dependencies(project, stage["id"], root)
            stage["upstream_changed"] = [
                s["title"]
                for s in project["stages"]
                if expected is not None
                and s["id"] in expected
                and current[s["id"]] != expected[s["id"]]
            ]
            stage["documents_changed"] = stage[
                "status"
            ] == "completed" and stage.get(
                "confirmed_forms",
            ) != self.form_fingerprint(
                root,
                project,
                stage["id"],
            )
            stage["missing_required"] = [
                m["name"]
                for m in stage["materials"]
                if m["required"] and m["status"] not in {"done", "na"}
            ]
            for material in stage["materials"]:
                for file in material["files"]:
                    approval = file.get("approval")
                    if approval:
                        file["approval_review"] = self.approval_review(
                            root,
                            project,
                            approval,
                        )
                        file.pop("approval", None)
                    source = (
                        self.path(root, project["id"])
                        / "files"
                        / file["storage"]
                    )
                    file["preview_url"] = "/api/files/preview/" + quote(
                        str(source),
                        safe="",
                    )
                    file.pop("storage", None)
        return result

    def approval_review(self, root, project, approval):
        review = copy.deepcopy(approval)
        if review.get("status") != "ready":
            return review
        current = values(project["fields"])
        case_id = project.get("form_id")
        if case_id:
            case = self.host.load(root, case_id)
            context = (
                case
                if case.get("stage_id") == "plan"
                else case.get("stage_contexts", {}).get("plan", {})
            )
            current.update(values(context.get("fields", [])))
        for field in review["fields"]:
            field["current"] = current.get(field["key"], "")
            field["changed"] = (
                field["current"].strip() != field["value"].strip()
            )
        review["change_count"] = sum(f["changed"] for f in review["fields"])
        return review

    def apply_approval(self, root, project, file_id, keys):
        file = next(
            (
                file
                for stage in project["stages"]
                for material in stage["materials"]
                for file in material["files"]
                if file["id"] == file_id
            ),
            None,
        )
        if not file or file.get("approval", {}).get("status") != "ready":
            raise HTTPException(404, "审批比对结果不存在")
        approved = {
            field["key"]: field["value"]
            for field in file["approval"]["fields"]
        }
        if set(keys) - set(approved):
            raise HTTPException(422, "所选审批字段无效")
        before = values(project["fields"])
        for field in project["fields"]:
            if field["key"] in keys:
                field["value"] = approved[field["key"]]
        if "project.name" in keys:
            project["name"] = approved["project.name"]
        shared_changed = before != values(project["fields"])
        case_id = project.get("form_id")
        if case_id:
            case = self.host.load(root, case_id)
            context = (
                case
                if case.get("stage_id") == "plan"
                else case.get("stage_contexts", {}).get("plan")
            )
            if context:
                snapshot = context.setdefault("project_snapshot", {})
                for field in context.get("fields", []):
                    if field["key"] in keys:
                        field["value"] = approved[field["key"]]
                        if field.get("scope") == "shared":
                            snapshot[field["key"]] = approved[field["key"]]
                context["review_token"] = None
                context["revision"] = (
                    int(context.get("revision", case["revision"])) + 1
                )
                case["revision"] += 1
                case["review_token"] = None
                self.host.save(root, case)
        if shared_changed:
            project["shared_revision"] += 1
        file["approval"]["applied_at"] = now()
        file["approval"]["applied_keys"] = sorted(set(keys))
        self.touch(project, self.stage(project, "plan"), root)
        self.save(root, project)
        return self.public(root, project)

    def add_case(self, root, project, stage_id, workbook=None):
        self.stage(project, stage_id)
        if workbook:
            try:
                data = self.host.read_excel(
                    self.host.upload(workbook, ".xlsx"),
                )
            except Exception as exc:
                if isinstance(exc, HTTPException):
                    raise
                raise HTTPException(422, f"Excel 无法读取：{exc}") from exc
        else:
            data = {
                "fields": (
                    self.host.builtin.blank_fields()
                    if stage_id == "negotiation"
                    else (
                        self.host.builtin.plan_fields(project["fields"])
                        if stage_id == "plan"
                        else (
                            self.host.builtin.selection_fields(
                                project["fields"],
                            )
                            if stage_id == "selection"
                            else (
                                self.host.builtin.acceptance_fields(
                                    project["fields"],
                                )
                                if stage_id == "acceptance"
                                else copy.deepcopy(project["fields"])
                            )
                        )
                    )
                ),
                "recipients": (
                    [{"name": "", "start": "", "end": ""} for _ in range(3)]
                    if stage_id == "negotiation"
                    else []
                ),
                "rules": [],
            }
        common = values(project["fields"])
        for field in data["fields"]:
            if not field["value"] and field["key"] in common:
                field["value"] = common[field["key"]]
        title = self.stage(project, stage_id)["title"]
        case = dict(
            data,
            id=uuid.uuid4().hex,
            name=title + "材料",
            revision=1,
            templates=[],
            generations=[],
            review_token=None,
            project_id=project["id"],
            stage_id=stage_id,
            project_snapshot={
                k: common[k] for k in common if k in values(data["fields"])
            },
        )
        try:
            self.host.attach_default_templates(root, case)
        except (ValueError, KeyError, OSError):
            case["template_notice"] = "默认模板未能载入，请重新选择 Word 文件"
        self.host.save(root, case)
        project["forms"].append({"case_id": case["id"], "stage_id": stage_id})
        self.touch(project, self.stage(project, stage_id), root)
        self.save(root, project)
        return self.host.public(case)

    @staticmethod
    def linked(project, case_id):
        link = next(
            (f for f in project["forms"] if f["case_id"] == case_id),
            None,
        )
        if not link:
            raise HTTPException(404, "该表单不属于此项目")
        return link


def register_project_routes(router, scope, lock, store):
    @router.get("/projects")
    def list_projects(root=Depends(scope)):
        with lock:
            projects = store.migrate(root)
            return [
                {
                    "id": p["id"],
                    "name": p["name"],
                    "code": p["code"],
                    "owner": p["owner"],
                    "updated_at": p["updated_at"],
                    "completed": sum(
                        s["status"] == "completed" for s in p["stages"]
                    ),
                    "forms": len(p["forms"]),
                }
                for p in sorted(
                    projects,
                    key=lambda p: p["updated_at"],
                    reverse=True,
                )
            ]

    @router.post("/projects")
    def create_project(body: ProjectBody, root=Depends(scope)):
        if not body.name.strip():
            raise HTTPException(422, "请填写项目名称")
        with lock:
            return store.public(root, store.create(root, body))

    @router.get("/projects/{project_id}")
    def get_project(project_id: str, root=Depends(scope)):
        with lock:
            return store.public(root, store.load(root, project_id))

    @router.put("/projects/{project_id}")
    def update_project(
        project_id: str,
        body: ProfileBody,
        root=Depends(scope),
    ):
        with lock:
            project = store.load(root, project_id)
            store.check(project, body.revision)
            if (
                not body.name.strip()
                or set(body.values) - set(values(project["fields"]))
                or any(len(v) > 20000 for v in body.values.values())
            ):
                raise HTTPException(422, "项目名称或共用字段无效")
            before = values(project["fields"])
            project.update(
                name=body.name.strip(),
                code=body.code.strip(),
                owner=body.owner.strip(),
            )
            for field in project["fields"]:
                if field["key"] in body.values:
                    field["value"] = body.values[field["key"]]
                if field["key"] == "project.name":
                    field["value"] = project["name"]
                if field["key"] == "project.code":
                    field["value"] = project["code"]
            if before != values(project["fields"]):
                project["shared_revision"] += 1
            project["revision"] += 1
            store.save(root, project)
            return store.public(root, project)


def register_stage_routes(router, scope, lock, store):
    @router.put("/projects/{project_id}/stages/{stage_id}")
    def update_stage(
        project_id: str,
        stage_id: str,
        body: StageBody,
        root=Depends(scope),
    ):
        with lock:
            project = store.load(root, project_id)
            store.check(project, body.revision)
            stage = store.stage(project, stage_id)
            if body.status not in STAGE_STATUS:
                raise HTTPException(422, "阶段状态无效")
            if body.due_date:
                try:
                    date.fromisoformat(body.due_date)
                except ValueError as exc:
                    raise HTTPException(422, "计划日期无效") from exc
            if body.system_url and not re.match(
                r"^https?://[^\s]+$",
                body.system_url,
                re.IGNORECASE,
            ):
                raise HTTPException(422, "系统地址须以 http:// 或 https:// 开头")
            missing = [
                m["name"]
                for m in stage["materials"]
                if m["required"] and m["status"] not in {"done", "na"}
            ]
            if body.status == "completed" and not body.confirmed:
                raise HTTPException(422, "请先确认本阶段已办结")
            if body.status == "completed" and missing:
                raise HTTPException(422, "必备材料尚未确认完成：" + "、".join(missing))
            if (
                body.status == "completed"
                and not body.notes.strip()
                and not any(m["status"] == "done" for m in stage["materials"])
            ):
                raise HTTPException(422, "请确认已完成的材料，或补充本阶段办结说明")
            store.touch(project, stage, root)
            stage.update(body.model_dump(exclude={"revision", "confirmed"}))
            stage["completed_at"] = (
                now() if body.status == "completed" else None
            )
            if body.status == "completed":
                stage["acknowledged"] = store.dependencies(
                    project,
                    stage_id,
                    root,
                )
                stage["confirmed_forms"] = store.form_fingerprint(
                    root,
                    project,
                    stage_id,
                )
            store.save(root, project)
            return store.public(root, project)

    @router.post("/projects/{project_id}/stages/{stage_id}/acknowledge")
    def acknowledge(
        project_id: str,
        stage_id: str,
        body: Revision,
        root=Depends(scope),
    ):
        with lock:
            project = store.load(root, project_id)
            store.check(project, body.revision)
            stage = store.stage(project, stage_id)
            stage["acknowledged"] = store.dependencies(project, stage_id, root)
            stage["confirmed_forms"] = store.form_fingerprint(
                root,
                project,
                stage_id,
            )
            project["revision"] += 1
            store.save(root, project)
            return store.public(root, project)


def register_material_routes(router, scope, lock, store):
    @router.post("/projects/{project_id}/stages/{stage_id}/materials")
    def add_material(
        project_id: str,
        stage_id: str,
        body: MaterialBody,
        root=Depends(scope),
    ):
        return write_material(project_id, stage_id, None, body, root)

    @router.put(
        "/projects/{project_id}/stages/{stage_id}/materials/{material_id}",
    )
    def write_material(
        project_id: str,
        stage_id: str,
        material_id: str,
        body: MaterialBody,
        root=Depends(scope),
    ):
        with lock:
            project = store.load(root, project_id)
            store.check(project, body.revision)
            stage = store.stage(project, stage_id)
            if body.status not in MATERIAL_STATUS or not body.name.strip():
                raise HTTPException(422, "材料名称或状态无效")
            if material_id:
                _, material = store.material(project, stage_id, material_id)
            else:
                if len(stage["materials"]) >= 200:
                    raise HTTPException(422, "单阶段最多 200 项材料")
                material = {"id": uuid.uuid4().hex, "files": []}
                stage["materials"].append(material)
            material.update(body.model_dump(exclude={"revision"}))
            material["name"] = material["name"].strip()
            store.touch(project, stage, root)
            store.save(root, project)
            return store.public(root, project)


def register_file_routes(router, scope, lock, store):
    host = store.host

    @router.post(
        "/projects/{project_id}/stages/{stage_id}"
        "/materials/{material_id}/files",
    )
    def upload_file(
        project_id: str,
        stage_id: str,
        material_id: str,
        revision: int = Form(...),
        file: UploadFile = File(...),
        root=Depends(scope),
    ):
        suffix = Path(file.filename or "").suffix.lower()
        if suffix not in EXTENSIONS:
            raise HTTPException(422, "支持 Office 文档、PDF、图片、ZIP、文本及 CSV 文件")
        content = host.upload(file, suffix)
        if not content:
            raise HTTPException(422, "不能上传空文件")
        with lock:
            project = store.load(root, project_id)
            store.check(project, revision)
            stage, material = store.material(project, stage_id, material_id)
            if len(material["files"]) >= 50:
                raise HTTPException(422, "单项材料最多保留 50 个文件版本")
            file_id = uuid.uuid4().hex
            storage = file_id + suffix
            path = store.path(root, project_id) / "files" / storage
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(content)
            material["files"].append(
                {
                    "id": file_id,
                    "name": host.name_only(Path(file.filename).stem)[:85]
                    + suffix,
                    "storage": storage,
                    "size": len(content),
                    "sha256": hashlib.sha256(content).hexdigest(),
                    "uploaded_at": now(),
                    "mime": mimetypes.guess_type(file.filename)[0]
                    or "application/octet-stream",
                },
            )
            if stage_id == "plan" and suffix == ".pdf":
                material["files"][-1]["approval"] = approval_from_pdf(content)
            material["status"] = "preparing"
            store.touch(project, stage, root)
            try:
                store.save(root, project)
            except OSError:
                path.unlink(missing_ok=True)
                raise
            return store.public(root, project)

    @router.post(
        "/projects/{project_id}/stages/plan/approval/apply",
    )
    def apply_plan_approval(
        project_id: str,
        body: ApprovalApplyBody,
        root=Depends(scope),
    ):
        with lock:
            project = store.load(root, project_id)
            store.check(project, body.revision)
            return store.apply_approval(root, project, body.file_id, body.keys)

    @router.get("/projects/{project_id}/files/{file_id}")
    def get_file(project_id: str, file_id: str, root=Depends(scope)):
        with lock:
            project = store.load(root, project_id)
            for stage in project["stages"]:
                for material in stage["materials"]:
                    for file in material["files"]:
                        if file["id"] == file_id:
                            source = (
                                store.path(root, project_id)
                                / "files"
                                / file["storage"]
                            )
                            if (
                                not source.is_file()
                                or hashlib.sha256(
                                    source.read_bytes(),
                                ).hexdigest()
                                != file["sha256"]
                            ):
                                raise HTTPException(422, "材料文件缺失或校验失败")
                            return FileResponse(
                                source,
                                filename=file["name"],
                                media_type=file["mime"],
                            )
            raise HTTPException(404, "材料文件不存在")


def register_form_routes(router, scope, lock, store):
    @router.post("/projects/{project_id}/form")
    def open_project_form(
        project_id: str,
        body: CaseBody,
        root=Depends(scope),
    ):
        with lock:
            project = store.load(root, project_id)
            store.check(project, body.revision)
            return store.open_form(root, project, body.stage_id)

    @router.post("/projects/{project_id}/form/import")
    def import_project_form(
        project_id: str,
        revision: int = Form(...),
        stage_id: str = Form(...),
        workbook: UploadFile = File(...),
        root=Depends(scope),
    ):
        with lock:
            project = store.load(root, project_id)
            store.check(project, revision)
            return store.open_form(root, project, stage_id, workbook)

    @router.post("/projects/{project_id}/cases")
    def create_form(project_id: str, body: CaseBody, root=Depends(scope)):
        with lock:
            project = store.load(root, project_id)
            store.check(project, body.revision)
            if project.get("form_id"):
                return store.open_form(root, project, body.stage_id)
            return store.add_case(root, project, body.stage_id)

    @router.post("/projects/{project_id}/cases/import")
    def import_form(
        project_id: str,
        revision: int = Form(...),
        stage_id: str = Form(...),
        workbook: UploadFile = File(...),
        root=Depends(scope),
    ):
        with lock:
            project = store.load(root, project_id)
            store.check(project, revision)
            if project.get("form_id"):
                return store.open_form(root, project, stage_id, workbook)
            return store.add_case(root, project, stage_id, workbook)

    @router.post("/projects/{project_id}/cases/{case_id}/sync")
    def sync_form(
        project_id: str,
        case_id: str,
        body: SyncBody,
        root=Depends(scope),
    ):
        with lock:
            project = store.load(root, project_id)
            store.check(project, body.revision)
            return store.sync(root, project, case_id, body)


def register_bundle_routes(router, scope, lock, store):
    host = store.host

    @router.get("/projects/{project_id}/bundle")
    def bundle(
        project_id: str,
        stage_id: str | None = None,
        root=Depends(scope),
    ):
        with lock:
            project = store.load(root, project_id)
            if stage_id:
                store.stage(project, stage_id)
            payload = store.archive(root, project, stage_id)
            suffix = (
                store.stage(project, stage_id)["title"]
                if stage_id
                else "全流程材料"
            )

        def chunks():
            try:
                while chunk := payload.read(1024 * 1024):
                    yield chunk
            finally:
                payload.close()

        return StreamingResponse(
            chunks(),
            media_type="application/zip",
            headers={
                "Content-Disposition": "attachment; filename*=UTF-8''"
                + quote(
                    host.name_only(project["name"]) + "_" + suffix + ".zip",
                ),
            },
        )


def register_routes(router, scope, lock, host):
    """Use the same authenticated user and Agent scope for every route."""
    store = ProjectStore(host)
    register_project_routes(router, scope, lock, store)
    register_stage_routes(router, scope, lock, store)
    register_material_routes(router, scope, lock, store)
    register_file_routes(router, scope, lock, store)
    register_form_routes(router, scope, lock, store)
    register_bundle_routes(router, scope, lock, store)
    return store
