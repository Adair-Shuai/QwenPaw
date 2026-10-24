"""Real staged engine contract. Run explicitly after staging OfficeCLI.

These tests never mock the CLI, and fail if the staged engine is missing.
Only disposable documents under pytest's temporary workspace are edited.
"""

import json
from pathlib import Path
import subprocess
import sys
import zipfile

import pytest

ROOT = Path(__file__).resolve().parents[2]
BINARY = (
    ROOT
    / "console/src-tauri/binaries/officecli"
    / ("officecli.exe" if sys.platform == "win32" else "officecli")
)


@pytest.fixture
def cli():
    assert BINARY.is_file(), "Stage the reviewed OfficeCLI engine first"

    def run(*args):
        result = subprocess.run(
            [str(BINARY), *map(str, args)],
            capture_output=True,
            encoding="utf-8",
            errors="replace",
            timeout=60,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        assert result.returncode == 0, result.stdout + result.stderr
        return result.stdout

    assert run("--version").strip() == "1.0.157"
    return run


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["docx", "xlsx", "pptx"])
async def test_tools_create_write_and_reopen_real_documents(
    cli, tmp_path, monkeypatch, extension
):
    import qwenpaw.agents.tools.office_tools as office

    monkeypatch.setattr(office, "_officecli_bin", lambda: str(BINARY))
    monkeypatch.setattr(
        office, "_resolve_workspace_path", lambda path: str(tmp_path / path)
    )
    monkeypatch.setattr(office, "get_current_workspace_dir", lambda: str(tmp_path))

    def unpack(chunk):
        return json.loads(chunk.content[0].text)

    name = "contract." + extension
    created = unpack(await office.office_create_document(name))
    assert created["ok"], created
    parent, kind, props = {
        "docx": (
            "/body",
            "paragraph",
            {"text": "Engine contract 中文\nSecond paragraph"},
        ),
        "xlsx": ("/", "sheet", {"name": "Contract"}),
        "pptx": ("/", "slide", {}),
    }[extension]
    added = unpack(await office.office_add_element(name, parent, kind, props))
    assert added["ok"], added
    if extension == "xlsx":
        added = unpack(
            await office.office_add_element(
                name, "/Contract", "cell", {"ref": "A1", "value": "Engine contract"}
            )
        )
        assert added["ok"], added
    elif extension == "pptx":
        added = unpack(
            await office.office_add_element(
                name, "/slide[1]", "shape", {"text": "Engine contract"}
            )
        )
        assert added["ok"], added

    path = tmp_path / name
    cli("close", path)
    with zipfile.ZipFile(path) as package:
        xml = "".join(
            package.read(item).decode("utf-8")
            for item in package.namelist()
            if item.endswith(".xml")
        )
    assert "Engine contract" in xml
    original = path.read_bytes()
    output = tmp_path / (extension + ".html")
    cli("view", path, "html", "-o", output)
    assert "Engine contract" in output.read_text(encoding="utf-8")
    assert path.read_bytes() == original, "Preview must not alter the source"


@pytest.mark.asyncio
async def test_csv_locale_import_persists_numbers(cli, tmp_path, monkeypatch):
    import qwenpaw.agents.tools.office_tools as office

    monkeypatch.setattr(office, "_officecli_bin", lambda: str(BINARY))
    monkeypatch.setattr(
        office, "_resolve_workspace_path", lambda path: str(tmp_path / path)
    )
    path = tmp_path / "locale.xlsx"
    csv = tmp_path / "values.csv"
    csv.write_text("name;value\npressure;12,5\n", encoding="utf-8")
    cli("create", path)
    chunk = await office.office_import_data(
        "locale.xlsx", "/Sheet1", str(csv), delimiter=";", decimal=","
    )
    result = json.loads(chunk.content[0].text)
    assert result["ok"], result
    cli("close", path)
    with zipfile.ZipFile(path) as package:
        sheet = package.read("xl/worksheets/sheet1.xml").decode()
    assert ">12.5<" in sheet


def test_word_comments_are_present_in_read_only_html(cli, tmp_path):
    from docx import Document

    doc = Document()
    run = doc.add_paragraph("Commented text").runs[0]
    doc.add_comment(run, "Review comment from UGSci", author="Reviewer")
    path = tmp_path / "comments.docx"
    doc.save(path)
    original = path.read_bytes()
    html = tmp_path / "comments.html"
    cli("view", path, "html", "-o", html)
    preview = html.read_text(encoding="utf-8")
    assert "Review comment from UGSci" in preview
    assert "Reviewer" in preview
    assert path.read_bytes() == original


def test_macro_payload_survives_edit_without_execution(cli, tmp_path):
    from docx import Document

    source = tmp_path / "source.docx"
    Document().save(source)
    macro = tmp_path / "preserved.docm"
    payload = b"inert VBA preservation regression fixture"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(macro, "w") as target:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                    b"application/vnd.ms-word.document.macroEnabled.main+xml",
                )
                data = data.replace(
                    b"</Types>",
                    b'<Override PartName="/word/vbaProject.bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>',
                )
            if item.filename == "word/_rels/document.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdMacro" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>',
                )
            target.writestr(item.filename, data)
        target.writestr("word/vbaProject.bin", payload)
    cli(
        "add",
        macro,
        "/body",
        "--type",
        "paragraph",
        "--prop",
        "text=Macro preservation edit",
    )
    cli("save", macro)
    cli("close", macro)
    with zipfile.ZipFile(macro) as result:
        assert result.read("word/vbaProject.bin") == payload
        assert b"Macro preservation edit" in result.read("word/document.xml")


@pytest.mark.asyncio
async def test_word_form_inspection_uses_real_engine(cli, tmp_path, monkeypatch):
    from docx import Document
    from docx.oxml import parse_xml
    import qwenpaw.agents.tools.office_tools as office

    doc = Document()
    doc.element.body.insert(
        0,
        parse_xml("""
      <w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
        <w:sdtPr><w:alias w:val="Company"/><w:tag w:val="company"/><w:id w:val="123"/><w:text/></w:sdtPr>
        <w:sdtContent><w:p><w:r><w:t>Initial company</w:t></w:r></w:p></w:sdtContent>
      </w:sdt>"""),
    )
    path = tmp_path / "form.docx"
    doc.save(path)
    monkeypatch.setattr(office, "_officecli_bin", lambda: str(BINARY))
    monkeypatch.setattr(
        office, "_resolve_workspace_path", lambda value: str(tmp_path / value)
    )
    chunk = await office.office_view_document("form.docx", "forms")
    result = json.loads(chunk.content[0].text)
    assert result["ok"], result
    assert "company" in json.dumps(result).lower()
    assert "Initial company" in json.dumps(result)
